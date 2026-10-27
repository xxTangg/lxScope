"""Task-owned RAG and citation integration tests using isolated fixtures.

These tests exercise the Task adapter contract. They do not claim live
knowledge-base or authenticated Workspace integration.
"""

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest import IsolatedAsyncioTestCase

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from examples.agent_service.task._citation_verifier import (
    verify_task_citations,
)
from examples.agent_service.task._executor import StepExecutionResult
from examples.agent_service.task._knowledge import (
    PLANNING_TEMPLATE_NO_MATCH_MESSAGE,
    PLANNING_TEMPLATE_SEARCH_QUERIES,
    TaskKnowledgeError,
    TaskKnowledgeGateway,
)
from examples.agent_service.task._models import (
    ArtifactConfig,
    ArtifactFormat,
    RunStatus,
    TaskContext,
    TaskInputReply,
    TaskPreflightRequest,
    TaskRunRequest,
    TaskToolSchema,
    UpdateTaskRequest,
)
from examples.agent_service.task._service import TaskService
from examples.agent_service.task._router import task_router
from examples.agent_service.task._store import TaskStore
from examples.agent_service.task._template_compiler import TaskTemplateCompiler
from examples.agent_service.task._template_models import (
    TemplateInstantiateRequest,
)
from examples.agent_service.task._template_registry import TemplateRegistry
from examples.agent_service.task._template_service import TaskTemplateService
from examples.agent_service.task._template_validation import (
    TaskValidationFailure,
)


_OWNER = "rag-test-owner"
_KB_ID = "kb-authorized"
_DOCUMENT_ID = "doc-planning"
_SOURCE_TEXT = (
    "第一章 项目概况；第二章 目标范围；"
    "第三章 数据现状；第四章 咨询工作方法。"
)
_TEMPLATE_CHUNK_1 = (
    "第五章 进度计划；第六章 费用约束；第七章 风险管理；"
    "第八章 成果与验收。策划报告写作要求：\n"
    + "\n".join(
        f"- B-{index:02d} 编写要求 {index}" for index in range(1, 8)
    )
)
_TEMPLATE_CHUNK_2 = "\n".join(
    f"- B-{index:02d} 写作要求 {index}" for index in range(8, 16)
)
_TEMPLATE_SOURCE = "02_策划报告编制模板_测试.docx"
_CHUNK_HASH = hashlib.sha256(_SOURCE_TEXT.encode("utf-8")).hexdigest()
_PROJECT_DOC_ID = "doc-project-material"
_PROJECT_SOURCE = "01_项目委托与需求说明_测试.docx"
_PROJECT_TEXT = (
    "项目建筑面积12000平方米；设备台账数量64台；"
    "项目预算800万元；建设周期12个月。"
)
_PROJECT_HASH = hashlib.sha256(_PROJECT_TEXT.encode("utf-8")).hexdigest()


def _view():
    return SimpleNamespace(
        id=_KB_ID,
        name="授权项目资料",
        description="",
        document_count=1,
        chunk_count=1,
        status_counts=SimpleNamespace(ready=1),
    )


def _document(document_id=_DOCUMENT_ID, source="项目策划依据.md"):
    return SimpleNamespace(
        id=document_id,
        status="ready",
        updated_at=datetime(2026, 10, 10, tzinfo=timezone.utc),
        data=SimpleNamespace(filename=source),
    )


def _search_result():
    return SimpleNamespace(
        document_id=_DOCUMENT_ID,
        score=0.93,
        chunk=SimpleNamespace(
            content=SimpleNamespace(text=_SOURCE_TEXT),
            source="项目策划依据.md",
            chunk_index=0,
            total_chunks=1,
        ),
    )


def _document_search_result(document_id, source, text, chunk_index):
    return SimpleNamespace(
        document_id=document_id,
        score=0.93 - chunk_index / 100,
        chunk=SimpleNamespace(
            content=SimpleNamespace(text=text),
            source=source,
            chunk_index=chunk_index,
            total_chunks=3,
        ),
    )


class _KnowledgeBaseFixture:
    """Minimal shape of the existing authorized KnowledgeBaseService API."""

    def __init__(self, *, owner: str = _OWNER, search_error: bool = False):
        self.owner = owner
        self.search_error = search_error
        self.view_calls = []
        self.search_calls = []
        self.document_calls = []
        self.view = _view()

    async def list_knowledge_base_views(
        self,
        user_id,
        *,
        knowledge_base_id=None,
        page=1,
        page_size=128,
    ):
        self.view_calls.append((user_id, knowledge_base_id, page, page_size))
        if user_id != self.owner or knowledge_base_id != _KB_ID:
            return [], 0
        return [self.view], 1

    async def search(self, user_id, knowledge_base_id, query, top_k=5):
        self.search_calls.append((user_id, knowledge_base_id, query, top_k))
        if self.search_error:
            raise RuntimeError("fixture search failure")
        if user_id != self.owner or knowledge_base_id != _KB_ID:
            return []
        return [_search_result()]

    async def get_document_status(
        self,
        user_id,
        knowledge_base_id,
        document_ids,
    ):
        self.document_calls.append(
            (user_id, knowledge_base_id, list(document_ids))
        )
        if user_id != self.owner or knowledge_base_id != _KB_ID:
            return []
        return [_document()] if _DOCUMENT_ID in document_ids else []


class _PlanningTemplateKnowledgeBaseFixture(_KnowledgeBaseFixture):
    """Authorized search fixture with a chunked template and distractors."""

    def __init__(self, *, mode="template"):
        super().__init__()
        self.mode = mode
        self.requirement_id = "doc-project-requirements"
        self.case_id = "doc-historical-case"
        self.hidden_template_id = "doc-hidden-template"
        self.project_id = _PROJECT_DOC_ID

    async def search(self, user_id, knowledge_base_id, query, top_k=5):
        self.search_calls.append((user_id, knowledge_base_id, query, top_k))
        if user_id != self.owner or knowledge_base_id != _KB_ID:
            return []
        if self.search_error:
            raise RuntimeError("fixture search failure")
        if self.mode == "unauthorized_template":
            return [
                _document_search_result(
                    self.hidden_template_id,
                    _TEMPLATE_SOURCE,
                    _SOURCE_TEXT,
                    0,
                ),
            ]
        if self.mode == "no_template":
            return [
                _document_search_result(
                    self.requirement_id,
                    "01_项目需求资料.docx",
                    "项目需求：需要编制项目策划报告并明确目标、费用和进度。",
                    0,
                ),
                _document_search_result(
                    self.case_id,
                    "03_历史案例.docx",
                    "历史案例的八周实施经验和指标仅供参考，不是策划模板。",
                    0,
                ),
            ]

        if (
            self.mode == "project_materials"
            and query not in PLANNING_TEMPLATE_SEARCH_QUERIES
        ):
            return [
                _document_search_result(
                    self.project_id,
                    _PROJECT_SOURCE,
                    _PROJECT_TEXT,
                    0,
                ),
            ]

        chunk_for_query = {
            PLANNING_TEMPLATE_SEARCH_QUERIES[0]: (0, _SOURCE_TEXT),
            PLANNING_TEMPLATE_SEARCH_QUERIES[1]: (1, _TEMPLATE_CHUNK_1),
            PLANNING_TEMPLATE_SEARCH_QUERIES[2]: (2, _TEMPLATE_CHUNK_2),
            PLANNING_TEMPLATE_SEARCH_QUERIES[3]: (1, _TEMPLATE_CHUNK_1),
        }
        chunk_index, text = chunk_for_query.get(query, (0, _SOURCE_TEXT))
        source = (
            ""
            if self.mode == "missing_source_metadata"
            else _TEMPLATE_SOURCE
        )
        return [
            _document_search_result(
                _DOCUMENT_ID,
                source,
                text,
                chunk_index,
            ),
        ]

    async def get_document_status(
        self,
        user_id,
        knowledge_base_id,
        document_ids,
    ):
        self.document_calls.append(
            (user_id, knowledge_base_id, list(document_ids)),
        )
        if user_id != self.owner or knowledge_base_id != _KB_ID:
            return []
        if self.mode == "unauthorized_template":
            return []
        if self.mode == "no_template":
            available = {
                self.requirement_id: _document(
                    self.requirement_id,
                    "01_项目需求资料.docx",
                ),
                self.case_id: _document(
                    self.case_id,
                    "03_历史案例.docx",
                ),
            }
            return [
                available[item]
                for item in document_ids
                if item in available
            ]
        if self.mode == "missing_source_metadata":
            return (
                [_document(_DOCUMENT_ID, None)]
                if _DOCUMENT_ID in document_ids
                else []
            )
        if self.mode == "project_materials":
            available = {
                _DOCUMENT_ID: _document(_DOCUMENT_ID, _TEMPLATE_SOURCE),
                self.project_id: _document(self.project_id, _PROJECT_SOURCE),
            }
            return [
                available[item]
                for item in document_ids
                if item in available
            ]
        return (
            [_document(_DOCUMENT_ID, _TEMPLATE_SOURCE)]
            if _DOCUMENT_ID in document_ids
            else []
        )


class _ArtifactExecutor:
    def __init__(self, *, project_materials=False, invalid_final=False):
        self.agent_calls = []
        self.files = {}
        self.tool_listing_calls = 0
        self.project_materials = project_materials
        self.invalid_final = invalid_final

    async def list_tools(self, context):
        del context
        self.tool_listing_calls += 1
        return [
            TaskToolSchema(
                name="planning.template_retrieve",
                input_schema={"type": "object"},
                is_read_only=True,
            ),
            TaskToolSchema(
                name="document.citation_verify",
                input_schema={"type": "object"},
                is_read_only=True,
            ),
        ]

    async def execute_node(self, *, node, previous_output, context):
        del context
        self.agent_calls.append(node.name)
        self.assert_citation_prompt(node.prompt)
        if self.project_materials:
            project_claim = "项目建筑面积12000平方米。[C1]"
            template_claim = "依据 B-01，方案明确写作要求。[C2]"
            pending_claim = "设备台账数量、预算及周期均待核实。"
            body = (
                "# 项目策划初稿\n\n"
                "## 项目概况\n"
                + project_claim
                + "\n\n## 模板要求\n"
                + template_claim
                + "\n\n## 待补资料\n"
                + pending_claim
                + "\n\n"
            )
            citations = [
                {
                    "citation_id": "C1",
                    "knowledge_base_id": _KB_ID,
                    "document_id": _PROJECT_DOC_ID,
                    "chunk_index": 0,
                    "content_sha256": _PROJECT_HASH,
                    "quote": _PROJECT_TEXT,
                },
                {
                    "citation_id": "C2",
                    "knowledge_base_id": _KB_ID,
                    "document_id": _DOCUMENT_ID,
                    "chunk_index": 1,
                    "content_sha256": hashlib.sha256(
                        _TEMPLATE_CHUNK_1.encode("utf-8")
                    ).hexdigest(),
                    "quote": _TEMPLATE_CHUNK_1,
                },
            ]
            claims = [
                {
                    "section": "项目概况",
                    "claim": project_claim,
                    "claim_type": "project_fact",
                    "citation_ids": ["C1"],
                },
                {
                    "section": "模板要求",
                    "claim": template_claim,
                    "claim_type": "template_rule",
                    "citation_ids": ["C2"],
                },
            ]
            unverified_items = [
                {
                    "section": "待补资料",
                    "claim": pending_claim,
                    "reason": "这些项目数据未由本次检索的项目资料支持。",
                }
            ]
        else:
            project_claim = ""
            template_claim = f"{_SOURCE_TEXT}[C1]"
            pending_claim = "项目面积、设备台账数量、预算及周期均待核实。"
            body = "# 项目策划初稿\n\n" f"{_SOURCE_TEXT}[C1]\n\n"
            citations = [
                {
                    "citation_id": "C1",
                    "knowledge_base_id": _KB_ID,
                    "document_id": _DOCUMENT_ID,
                    "chunk_index": 0,
                    "content_sha256": _CHUNK_HASH,
                    "quote": _SOURCE_TEXT,
                }
            ]
            claims = [
                {
                    "section": "项目策划初稿",
                    "claim": template_claim,
                    "claim_type": "template_rule",
                    "citation_ids": ["C1"],
                }
            ]
            unverified_items = [
                {
                    "section": "项目策划初稿",
                    "claim": pending_claim,
                    "reason": "本次运行未检索到项目委托资料。",
                }
            ]
        if node.name == "整理策划输出":
            if pending_claim not in body:
                body += pending_claim + "\n\n"
            if self.invalid_final:
                invalid_claim = "项目建筑面积12000平方米。[C1]"
                body += invalid_claim + "\n\n"
                claims = [
                    {
                        "section": "项目策划初稿",
                        "claim": invalid_claim,
                        "claim_type": "project_fact",
                        "citation_ids": ["C1"],
                    }
                ]
            claims_block = (
                "```task-evidence-claims\n"
                + json.dumps(
                    {
                        "claims": claims,
                        "unverified_items": unverified_items,
                    },
                    ensure_ascii=False,
                    separators=(",", ":"),
                )
                + "\n```\n\n"
            )
            body += claims_block
        output = (
            body
            + "```task-citations\n"
            + json.dumps(
                {"citations": citations},
                ensure_ascii=False,
                separators=(",", ":"),
            )
            + "\n```"
        )
        return StepExecutionResult(text=output)

    @staticmethod
    def assert_citation_prompt(prompt):
        assert "task-citations" in prompt
        assert "逐字摘录" in prompt

    async def verify_node(self, **kwargs):
        del kwargs
        from examples.agent_service.task._models import StepVerdict

        return StepVerdict(passed=True)

    async def preflight_artifact_writer(self, *, context):
        assert context.user_id == _OWNER
        return True, "fixture Workspace boundary available"

    async def write_artifact(self, *, path, data, context):
        assert context.user_id == _OWNER
        self.files[path] = data
        return len(data)

    async def read_artifact(self, *, path, context):
        assert context.user_id == _OWNER
        return self.files[path]


class _OccurrenceProtocolArtifactExecutor(_ArtifactExecutor):
    async def execute_node(self, *, node, previous_output, context):
        if node.name != "整理策划输出":
            return await super().execute_node(
                node=node,
                previous_output=previous_output,
                context=context,
            )

        retrievals = []
        for item in context.history:
            structured = item.get("structured_result")
            if not isinstance(structured, dict):
                try:
                    structured = json.loads(item.get("output", ""))
                except (TypeError, json.JSONDecodeError):
                    continue
            if isinstance(structured, dict) and isinstance(
                structured.get("evidence"), list
            ):
                retrievals.append(structured)
        evidence = [
            source
            for retrieval in retrievals
            for source in retrieval.get("evidence", [])
            if isinstance(source, dict)
        ]
        project = next(
            item for item in evidence
            if item.get("evidence_type") == "project_material"
        )
        template = next(
            item for item in evidence
            if item.get("evidence_type") == "planning_template"
            and "B-01" in item.get("source_text", "")
        )
        project_quote = "项目建筑面积12000平方米"
        template_quote = template["source_text"]
        project_text = "项目建筑面积12000平方米。"
        rule_text = "依据 B-01 编写要求 1。"
        analysis_text = "分析应先补齐计量边界信息，再比较设备运行时段。"
        recommendation_text = "建议确认设备台账边界后再计算节能机会。"
        pending_text = "复核人待确认。"
        supplement_text = "表计边界图待补充。"
        body = (
            "# 项目策划草稿\n\n"
            "## 项目事实\n"
            f"{project_text}[C1] "
            "<!-- task-claim:CL-area:O-area -->\n\n"
            "## 模板依据\n"
            f"{rule_text}[C2] "
            "<!-- task-claim:CL-rule:O-rule -->\n\n"
            "## 分析和建议\n"
            f"{analysis_text} <!-- task-claim:CL-analysis:O-analysis -->\n\n"
            f"{recommendation_text} <!-- task-claim:CL-recommendation:O-recommendation -->\n\n"
            "## 待核实\n"
            f"{pending_text} <!-- task-claim:CL-reviewer:O-reviewer -->\n\n"
            f"{supplement_text} <!-- task-claim:CL-reviewer:O-supplement -->"
        )
        payload = {
            "claims": [
                {
                    "claim_id": "CL-area",
                    "kind": "fact",
                    "anchor": "CL-area",
                    "text": project_text,
                    "citation_ids": ["C1"],
                    "occurrences": [
                        {
                            "occurrence_id": "O-area",
                            "text": project_text,
                            "citation_ids": ["C1"],
                        }
                    ],
                },
                {
                    "claim_id": "CL-rule",
                    "kind": "rule",
                    "anchor": "CL-rule",
                    "text": rule_text,
                    "citation_ids": ["C2"],
                    "occurrences": [
                        {
                            "occurrence_id": "O-rule",
                            "text": rule_text,
                            "citation_ids": ["C2"],
                        }
                    ],
                },
                {
                    "claim_id": "CL-analysis",
                    "kind": "analysis",
                    "anchor": "CL-analysis",
                    "text": analysis_text,
                    "citation_ids": [],
                    "occurrences": [
                        {
                            "occurrence_id": "O-analysis",
                            "text": analysis_text,
                            "citation_ids": [],
                        }
                    ],
                },
                {
                    "claim_id": "CL-recommendation",
                    "kind": "recommendation",
                    "anchor": "CL-recommendation",
                    "text": recommendation_text,
                    "citation_ids": [],
                    "occurrences": [
                        {
                            "occurrence_id": "O-recommendation",
                            "text": recommendation_text,
                            "citation_ids": [],
                        }
                    ],
                },
                {
                    "claim_id": "CL-reviewer",
                    "kind": "unverified",
                    "anchor": "CL-reviewer",
                    "text": f"{pending_text} {supplement_text}",
                    "citation_ids": [],
                    "occurrences": [
                        {
                            "occurrence_id": "O-reviewer",
                            "text": pending_text,
                            "citation_ids": [],
                        },
                        {
                            "occurrence_id": "O-supplement",
                            "text": supplement_text,
                            "citation_ids": [],
                        },
                    ],
                },
            ],
            "bindings": [
                {
                    "claim_id": "CL-area",
                    "occurrence_id": "O-area",
                    "citation_id": "C1",
                    "evidence_id": project["evidence_id"],
                    "quote": project_quote,
                },
                {
                    "claim_id": "CL-rule",
                    "occurrence_id": "O-rule",
                    "citation_id": "C2",
                    "evidence_id": template["evidence_id"],
                    "quote": template_quote,
                },
            ],
            "unverified_items": [
                {
                    "unverified_id": "UV-reviewer",
                    "claim_id": "CL-reviewer",
                    "occurrence_id": "O-reviewer",
                    "status": "pending_confirmation",
                    "reason": "项目方尚未指定复核人。",
                    "expected_evidence_types": ["project_material"],
                },
                {
                    "unverified_id": "UV-meter-boundary",
                    "claim_id": "CL-reviewer",
                    "occurrence_id": "O-supplement",
                    "status": "pending_supplement",
                    "reason": "尚未取得表计边界图。",
                    "expected_evidence_types": ["project_material"],
                },
            ],
        }
        citations = {
            "citations": [
                {
                    "citation_id": "C1",
                    "evidence_id": project["evidence_id"],
                    "quote": project_quote,
                },
                {
                    "citation_id": "C2",
                    "evidence_id": template["evidence_id"],
                    "quote": template_quote,
                },
            ]
        }
        output = (
            body
            + "\n\n```task-claim-evidence\n"
            + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
            + "\n```\n\n```task-citations\n"
            + json.dumps(citations, ensure_ascii=False, separators=(",", ":"))
            + "\n```"
        )
        return StepExecutionResult(text=output)


class _IncompleteOccurrenceProtocolArtifactExecutor(
    _OccurrenceProtocolArtifactExecutor
):
    async def execute_node(self, *, node, previous_output, context):
        result = await super().execute_node(
            node=node,
            previous_output=previous_output,
            context=context,
        )
        if node.name != "整理策划输出":
            return result
        claim_fence = "\n```\n\n```task-citations\n"
        assert claim_fence in result.text
        self.raw_final_output = result.text.replace(
            claim_fence,
            "\n\n```task-citations\n",
            1,
        )
        return StepExecutionResult(
            text=self.raw_final_output,
        )


class _NeverPlanner:
    def __init__(self):
        self.calls = 0

    async def plan(self, **kwargs):
        del kwargs
        self.calls += 1
        raise AssertionError("A standard template must not call the Planner.")


class TaskProjectPlanningRagTest(IsolatedAsyncioTestCase):
    async def test_demo_planning_template_is_valid_across_rag_fragments(self):
        demo_path = (
            Path(__file__).resolve().parents[1]
            / "examples"
            / "agent_service"
            / "task"
            / "docs"
            / "p7b-demo-data"
            / "knowledge-sources"
            / "DEMO-SOURCE-002-planning-template.md"
        )
        original = demo_path.read_text(encoding="utf-8")
        fourth_chapter = original.index("第四章")
        rule_section = original.index("## B-01")
        eighth_rule = original.index("- B-08")
        chunks = [
            original[:fourth_chapter],
            original[fourth_chapter:rule_section],
            original[rule_section:eighth_rule],
            original[eighth_rule:],
        ]
        source = "DEMO-SOURCE-002-planning-template.md"
        document_id = "demo-authorized-planning-template"

        class DemoTemplateKnowledgeBase(_KnowledgeBaseFixture):
            async def search(
                self,
                user_id,
                knowledge_base_id,
                query,
                top_k=5,
            ):
                self.search_calls.append(
                    (user_id, knowledge_base_id, query, top_k),
                )
                if user_id != self.owner or knowledge_base_id != _KB_ID:
                    return []
                chunk_index = PLANNING_TEMPLATE_SEARCH_QUERIES.index(query)
                if getattr(self, "missing_chunk_index", None) == chunk_index:
                    return []
                return [
                    _document_search_result(
                        document_id,
                        source,
                        chunks[chunk_index],
                        chunk_index,
                    ),
                ]

            async def get_document_status(
                self,
                user_id,
                knowledge_base_id,
                document_ids,
            ):
                self.document_calls.append(
                    (user_id, knowledge_base_id, list(document_ids)),
                )
                if (
                    user_id != self.owner
                    or knowledge_base_id != _KB_ID
                    or document_id not in document_ids
                ):
                    return []
                return [_document(document_id, source)]

        knowledge_service = DemoTemplateKnowledgeBase()
        gateway = TaskKnowledgeGateway(knowledge_service, SimpleNamespace())

        result = await gateway.retrieve_planning_template_evidence(
            _OWNER,
            [_KB_ID],
        )

        self.assertEqual(len(result["template_matches"]), 1)
        match = result["template_matches"][0]
        self.assertEqual(match["source"], source)
        self.assertEqual(match["chapter_numbers"], list(range(1, 9)))
        self.assertEqual(
            match["writing_rule_ids"],
            [f"B-{index:02d}" for index in range(1, 16)],
        )
        self.assertTrue(match["has_outline_marker"])
        self.assertTrue(match["has_writing_requirements"])
        self.assertEqual(
            [item["chunk_index"] for item in result["evidence"]],
            [0, 1, 2, 3],
        )
        self.assertEqual(
            "".join(item["source_text"] for item in result["evidence"]),
            original,
        )
        self.assertTrue(
            all(item["document_id"] == document_id for item in result["evidence"]),
        )

        knowledge_service.missing_chunk_index = 3
        with self.assertRaises(TaskKnowledgeError) as context:
            await gateway.retrieve_planning_template_evidence(
                _OWNER,
                [_KB_ID],
            )
        rejected = context.exception.details["rejected_documents"][0]
        self.assertEqual(
            rejected["reason"],
            "授权原文未完整覆盖 B-01～B-15 编写要求。",
        )
        self.assertEqual(
            rejected["writing_rule_ids"],
            [f"B-{index:02d}" for index in range(1, 8)],
        )

    async def test_template_queries_are_bounded_and_merge_template_chunks(
        self,
    ):
        service = _PlanningTemplateKnowledgeBaseFixture()
        gateway = TaskKnowledgeGateway(service, SimpleNamespace())

        result = await gateway.retrieve_planning_template_evidence(
            _OWNER,
            [_KB_ID],
        )

        self.assertEqual(
            result["queries"],
            list(PLANNING_TEMPLATE_SEARCH_QUERIES),
        )
        self.assertEqual(
            [call[2] for call in service.search_calls],
            list(PLANNING_TEMPLATE_SEARCH_QUERIES),
        )
        self.assertTrue(
            all("形成项目策划初稿" not in call[2] for call in service.search_calls)
        )
        self.assertEqual(len(result["template_matches"]), 1)
        match = result["template_matches"][0]
        self.assertEqual(match["source"], _TEMPLATE_SOURCE)
        self.assertEqual(match["chapter_numbers"], list(range(1, 9)))
        self.assertEqual(
            match["writing_rule_ids"],
            [f"B-{index:02d}" for index in range(1, 16)],
        )
        self.assertEqual(
            [item["chunk_index"] for item in result["evidence"]],
            [0, 1, 2],
        )
        for item in result["evidence"]:
            self.assertEqual(item["knowledge_base_name"], "授权项目资料")
            self.assertEqual(item["document_id"], _DOCUMENT_ID)
            self.assertTrue(item["source_text"])
            self.assertEqual(len(item["content_sha256"]), 64)
            self.assertEqual(item["source"], _TEMPLATE_SOURCE)
        self.assertEqual(result["evidence"][0]["source_text"], _SOURCE_TEXT)
        self.assertEqual(result["evidence"][0]["content_sha256"], _CHUNK_HASH)

    async def test_project_requirements_and_cases_do_not_count_as_template(
        self,
    ):
        service = _PlanningTemplateKnowledgeBaseFixture(mode="no_template")
        gateway = TaskKnowledgeGateway(service, SimpleNamespace())

        with self.assertRaises(TaskKnowledgeError) as context:
            await gateway.retrieve_planning_template_evidence(
                _OWNER,
                [_KB_ID],
            )

        error = context.exception
        self.assertTrue(
            error.message.startswith(PLANNING_TEMPLATE_NO_MATCH_MESSAGE),
        )
        self.assertEqual(error.code, "no_matching_template")
        self.assertEqual(
            error.details["queries"],
            list(PLANNING_TEMPLATE_SEARCH_QUERIES),
        )
        self.assertEqual(
            {item["source"] for item in error.details["hit_documents"]},
            {"01_项目需求资料.docx", "03_历史案例.docx"},
        )
        self.assertIn("未满足策划模板来源", error.details["failure_reason"])
        self.assertIn("定向检索词", error.message)

    async def test_unauthorized_source_evidence_is_rejected(
        self,
    ):
        unauthorized = TaskKnowledgeGateway(
            _PlanningTemplateKnowledgeBaseFixture(
                mode="unauthorized_template",
            ),
            SimpleNamespace(),
        )
        with self.assertRaises(TaskKnowledgeError) as unauthorized_error:
            await unauthorized.retrieve_planning_template_evidence(
                _OWNER,
                [_KB_ID],
            )
        self.assertEqual(
            unauthorized_error.exception.code,
            "no_matching_template",
        )
        self.assertEqual(
            unauthorized_error.exception.details["hit_documents"],
            [],
        )

        missing_source = TaskKnowledgeGateway(
            _PlanningTemplateKnowledgeBaseFixture(
                mode="missing_source_metadata",
            ),
            SimpleNamespace(),
        )
        with self.assertRaises(TaskKnowledgeError) as metadata_error:
            await missing_source.retrieve_planning_template_evidence(
                _OWNER,
                [_KB_ID],
            )
        self.assertEqual(metadata_error.exception.code, "no_matching_template")
        self.assertTrue(
            any(
                item["reason"] == "检索原文缺少完整且真实的来源字段。"
                for item in metadata_error.exception.details[
                    "rejected_documents"
                ]
            )
        )

    async def test_retrieval_preserves_scope_and_source_trace(self):
        service = _KnowledgeBaseFixture()
        gateway = TaskKnowledgeGateway(service, SimpleNamespace())

        result = await gateway.retrieve_task_evidence(
            _OWNER,
            [_KB_ID],
            "项目策划目标",
        )

        evidence = result["evidence"][0]
        self.assertEqual(service.search_calls[0][0], _OWNER)
        self.assertEqual(service.document_calls[0][0], _OWNER)
        self.assertEqual(evidence["document_id"], _DOCUMENT_ID)
        self.assertEqual(evidence["knowledge_base_id"], _KB_ID)
        self.assertEqual(evidence["source"], "项目策划依据.md")
        self.assertEqual(evidence["content_sha256"], _CHUNK_HASH)
        self.assertEqual(evidence["source_text"], _SOURCE_TEXT)
        self.assertEqual(evidence["source_status"], "ready")

    async def test_inaccessible_knowledge_base_is_rejected_before_search(self):
        service = _KnowledgeBaseFixture()
        gateway = TaskKnowledgeGateway(service, SimpleNamespace())

        with self.assertRaisesRegex(RuntimeError, "不存在或当前用户无权"):
            await gateway.retrieve_task_evidence(
                "another-user",
                [_KB_ID],
                "项目策划目标",
            )

        self.assertEqual(service.search_calls, [])

    async def test_search_failure_and_empty_evidence_fail_closed(self):
        failed = TaskKnowledgeGateway(
            _KnowledgeBaseFixture(search_error=True),
            SimpleNamespace(),
        )
        with self.assertRaisesRegex(RuntimeError, "授权知识库检索失败"):
            await failed.retrieve_task_evidence(_OWNER, [_KB_ID], "项目策划目标")

        empty_service = _KnowledgeBaseFixture()
        empty_service.search = _empty_search
        empty = TaskKnowledgeGateway(empty_service, SimpleNamespace())
        with self.assertRaisesRegex(RuntimeError, "没有检索到可引用"):
            await empty.retrieve_task_evidence(_OWNER, [_KB_ID], "项目策划目标")

    async def test_citation_verifier_checks_identity_and_exact_quote(self):
        gateway = TaskKnowledgeGateway(
            _KnowledgeBaseFixture(),
            SimpleNamespace(),
        )
        retrieval = await gateway.retrieve_task_evidence(
            _OWNER,
            [_KB_ID],
            "项目策划目标",
        )
        document = _citation_document()

        verified = verify_task_citations(document, retrieval["evidence"])
        self.assertTrue(verified["passed"])
        self.assertEqual(
            verified["verified_citations"][0]["status"],
            "verified",
        )

        forged = document.replace(
            f'"quote":"{_SOURCE_TEXT}"',
            '"quote":"没有出现在资料中的内容"',
        )
        rejected = verify_task_citations(forged, retrieval["evidence"])
        self.assertFalse(rejected["passed"])
        self.assertTrue(
            any("摘录与保存的原文" in issue for issue in rejected["issues"])
        )

        unlisted_marker = document.replace("[C1]", "[C2]", 1)
        rejected_marker = verify_task_citations(
            unlisted_marker,
            retrieval["evidence"],
        )
        self.assertFalse(rejected_marker["passed"])
        self.assertTrue(
            any(
                "未列入核验清单" in issue
                for issue in rejected_marker["issues"]
            )
        )

        trailing_text = document + "\n新增未核验结论。"
        rejected_trailing = verify_task_citations(
            trailing_text,
            retrieval["evidence"],
        )
        self.assertFalse(rejected_trailing["passed"])
        self.assertIn(
            "task-citations 区块之后",
            rejected_trailing["issues"][0],
        )

    async def test_preflight_searches_and_confirms_configured_artifact(self):
        executor = _ArtifactExecutor()
        task_service = TaskService(
            TaskStore(),
            agentscope_executor=executor,
            preview_executor=executor,
            artifact_writer=executor,
        )
        template_service = TaskTemplateService(
            TemplateRegistry(),
            task_service,
            compiler=TaskTemplateCompiler(),
        )
        task = await template_service.instantiate(
            _OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version="2.0.0",
                inputs={
                    "project_name": "测试项目",
                    "project_brief": "项目简报请重复引用这段内容：形成项目策划初稿。",
                    "professional_fields": ["项目管理"],
                    "source_material_ids": [_DOCUMENT_ID],
                },
                source_context=TaskContext(
                    agent_id="agent",
                    session_id="session",
                ),
            ),
        )
        await task_service.update_task(
            _OWNER,
            task.id,
            UpdateTaskRequest(knowledge_base_ids=[_KB_ID]),
        )
        gateway = TaskKnowledgeGateway(
            _PlanningTemplateKnowledgeBaseFixture(),
            SimpleNamespace(),
        )

        preflight = await task_service.preflight_task(
            _OWNER,
            task.id,
            TaskPreflightRequest(input="项目策划初稿"),
            gateway,
        )

        checks = {
            check.capability_key: check
            for check in preflight.capability_checks
        }
        self.assertEqual(
            checks["planning.template_retrieve"].status,
            "available",
        )
        self.assertIn(
            "已匹配策划模板",
            checks["planning.template_retrieve"].message,
        )
        self.assertFalse(
            checks["planning.template_retrieve"].business_data_verified,
        )
        self.assertIn(
            "项目事实来源将在输入节点后单独检索",
            checks["planning.template_retrieve"].message,
        )
        self.assertEqual(
            [call[2] for call in gateway._service.search_calls],
            list(PLANNING_TEMPLATE_SEARCH_QUERIES),
        )
        self.assertTrue(
            all(
                "项目简报请重复引用" not in call[2]
                for call in gateway._service.search_calls
            ),
        )
        self.assertEqual(
            checks["document.citation_verify"].status,
            "available",
        )
        self.assertFalse(
            checks["document.citation_verify"].business_data_verified,
        )
        self.assertEqual(checks["document.render"].status, "available")
        self.assertNotIn("document.render", preflight.missing_capabilities)
        self.assertTrue(preflight.ready)
        await task_service.shutdown()

    async def test_preflight_rejects_only_project_doc_hits(
        self,
    ):
        executor = _ArtifactExecutor()
        task_service = TaskService(
            TaskStore(),
            agentscope_executor=executor,
            preview_executor=executor,
            artifact_writer=executor,
        )
        template_service = TaskTemplateService(
            TemplateRegistry(),
            task_service,
            compiler=TaskTemplateCompiler(),
        )
        task = await template_service.instantiate(
            _OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version="2.0.0",
                inputs={
                    "project_name": "测试项目",
                    "project_brief": "项目需求应优先形成完整策划报告。",
                    "professional_fields": ["项目管理"],
                },
            ),
        )
        await task_service.update_task(
            _OWNER,
            task.id,
            UpdateTaskRequest(knowledge_base_ids=[_KB_ID]),
        )
        knowledge_service = _PlanningTemplateKnowledgeBaseFixture(
            mode="no_template",
        )
        gateway = TaskKnowledgeGateway(knowledge_service, SimpleNamespace())

        preflight = await task_service.preflight_task(
            _OWNER,
            task.id,
            TaskPreflightRequest(input="项目需求应优先形成完整策划报告。"),
            gateway,
        )

        checks = {
            check.capability_key: check
            for check in preflight.capability_checks
        }
        check = checks["planning.template_retrieve"]
        self.assertEqual(check.status, "missing")
        self.assertFalse(check.business_data_verified)
        self.assertTrue(
            check.message.startswith(PLANNING_TEMPLATE_NO_MATCH_MESSAGE),
        )
        self.assertIn("01_项目需求资料.docx", check.message)
        self.assertIn("03_历史案例.docx", check.message)
        self.assertIn("定向检索词", check.message)
        self.assertFalse(preflight.ready)
        self.assertEqual(
            [call[2] for call in knowledge_service.search_calls],
            list(PLANNING_TEMPLATE_SEARCH_QUERIES),
        )
        with self.assertRaises(TaskValidationFailure):
            await task_service.create_run(
                _OWNER,
                task.id,
                TaskRunRequest(),
                gateway,
            )
        self.assertEqual(
            [call[2] for call in knowledge_service.search_calls[-4:]],
            list(PLANNING_TEMPLATE_SEARCH_QUERIES),
        )
        self.assertEqual(executor.agent_calls, [])
        self.assertEqual(executor.files, {})
        await task_service.shutdown()

    async def test_preflight_route_passes_existing_knowledge_service(self):
        executor = _ArtifactExecutor()
        task_service = TaskService(
            TaskStore(),
            agentscope_executor=executor,
            preview_executor=executor,
            artifact_writer=executor,
        )
        template_service = TaskTemplateService(
            TemplateRegistry(),
            task_service,
            compiler=TaskTemplateCompiler(),
        )
        task = await template_service.instantiate(
            _OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version="2.0.0",
                inputs={
                    "project_name": "测试项目",
                    "project_brief": "形成项目策划初稿",
                    "professional_fields": ["项目管理"],
                    "source_material_ids": [_DOCUMENT_ID],
                },
            ),
        )
        await task_service.update_task(
            _OWNER,
            task.id,
            UpdateTaskRequest(knowledge_base_ids=[_KB_ID]),
        )
        app = FastAPI()
        app.include_router(task_router)
        app.state.task_service = task_service
        app.state.knowledge_base_service = (
            _PlanningTemplateKnowledgeBaseFixture()
        )

        with TestClient(app) as client:
            response = client.post(
                f"/tasks/{task.id}/preflight",
                headers={"X-User-ID": _OWNER},
                json={"input": "项目策划初稿"},
            )

        self.assertEqual(response.status_code, 200, response.text)
        checks = {
            check["capability_key"]: check
            for check in response.json()["capability_checks"]
        }
        self.assertEqual(
            checks["planning.template_retrieve"]["status"],
            "available",
        )
        await task_service.shutdown()

    async def test_s02_custom_copy_runs_rag_review_and_word_writer(self):
        executor = _ArtifactExecutor()
        planner = _NeverPlanner()
        store = TaskStore()
        task_service = TaskService(
            store,
            agentscope_executor=executor,
            preview_executor=executor,
            artifact_writer=executor,
            planner=planner,
        )
        template_service = TaskTemplateService(
            TemplateRegistry(),
            task_service,
            compiler=TaskTemplateCompiler(),
        )
        task = await template_service.instantiate(
            _OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version="2.0.0",
                inputs={
                    "project_name": "测试项目",
                    "project_brief": "项目简报请重复引用这段内容：形成项目策划初稿。",
                    "professional_fields": ["项目管理"],
                    "source_material_ids": [_DOCUMENT_ID],
                },
                source_context=TaskContext(
                    agent_id="agent",
                    session_id="session",
                ),
            ),
        )
        self.assertEqual(planner.calls, 0)
        self.assertEqual(executor.agent_calls, [])
        self.assertEqual(executor.files, {})
        copied = await task_service.create_custom_flow_copy(_OWNER, task.id)
        self.assertIsNotNone(copied)
        final_node = copied.nodes[-1].model_copy(
            update={
                "artifact": ArtifactConfig(
                    format=ArtifactFormat.DOCX,
                    filename="项目策划初稿.docx",
                ),
            },
        )
        custom = await task_service.update_task(
            _OWNER,
            copied.id,
            UpdateTaskRequest(
                knowledge_base_ids=[_KB_ID],
                nodes=[*copied.nodes[:-1], final_node],
            ),
        )
        knowledge_service = _PlanningTemplateKnowledgeBaseFixture()
        gateway = TaskKnowledgeGateway(knowledge_service, SimpleNamespace())

        run = await task_service.create_run(
            _OWNER,
            custom.id,
            TaskRunRequest(),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)
        self.assertEqual(run.status, RunStatus.AWAITING_INPUT)
        run = await task_service.answer_run(
            _OWNER,
            run.id,
            TaskInputReply(
                request_id=run.pending_input.id,
                answers={
                    "project_name": "测试项目",
                    "project_brief": "形成项目策划初稿",
                },
            ),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)
        self.assertEqual(run.status, RunStatus.AWAITING_INPUT)
        run = await task_service.answer_run(
            _OWNER,
            run.id,
            TaskInputReply(
                request_id=run.pending_input.id,
                answers={"planning_feedback": "无补充意见"},
            ),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)

        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        self.assertEqual(len(executor.agent_calls), 3)
        self.assertEqual(executor.tool_listing_calls, 0)
        self.assertEqual(planner.calls, 0)
        self.assertEqual(len(executor.files), 1)
        artifact = run.artifacts[0]
        self.assertEqual(artifact.format, ArtifactFormat.DOCX)
        docx_file = executor.files[artifact.path]
        self.assertTrue(docx_file.startswith(b"PK"))
        document = Document(BytesIO(docx_file))
        document_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
        )
        self.assertIn(_SOURCE_TEXT, document_text)
        self.assertIn("task-citations", document_text)
        self.assertIn("[C1]", run.final_output)
        self.assertEqual(
            run.node_runs[4].metadata["citation_protocol_status"],
            "verified",
        )
        self.assertTrue(
            run.node_runs[4].metadata["citation_verification"]["passed"]
        )
        self.assertTrue(
            run.node_runs[-1].metadata["final_citation_verification"]["passed"]
        )
        self.assertTrue(run.node_runs[-1].metadata["evidence_ledger"]["passed"])
        self.assertFalse(
            any(
                item["error_code"] == "legacy_compatibility_validation"
                for item in run.node_runs[-1].metadata["evidence_ledger"].get(
                    "all_diagnostics", []
                )
            )
        )
        retrieval_result = run.node_runs[1].result
        self.assertTrue(
            all(
                item["evidence_id"].startswith("EV-")
                and item["authorization_subject"] == _OWNER
                for item in retrieval_result["evidence"]
            )
        )
        self.assertIn("evidence_id", run.node_runs[1].output)
        self.assertNotIn("authorization_subject", run.node_runs[1].output)
        self.assertEqual(run.source_material_ids, [_DOCUMENT_ID])
        self.assertEqual(
            retrieval_result["queries"],
            list(PLANNING_TEMPLATE_SEARCH_QUERIES),
        )
        self.assertEqual(
            retrieval_result["template_matches"][0]["source"],
            _TEMPLATE_SOURCE,
        )
        self.assertTrue(
            all(
                "项目简报请重复引用" not in call[2]
                for call in knowledge_service.search_calls
            ),
        )
        self.assertEqual(planner.calls, 0)
        await task_service.shutdown()

    async def test_input_selected_project_material_is_saved_in_run_and_ledger(
        self,
    ):
        executor = _ArtifactExecutor(project_materials=True)
        planner = _NeverPlanner()
        task_service = TaskService(
            TaskStore(),
            agentscope_executor=executor,
            preview_executor=executor,
            artifact_writer=executor,
            planner=planner,
        )
        template_service = TaskTemplateService(
            TemplateRegistry(),
            task_service,
            compiler=TaskTemplateCompiler(),
        )
        task = await template_service.instantiate(
            _OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version="2.0.0",
                inputs={
                    "project_name": "示范工程",
                    "project_brief": "依据项目委托资料形成策划报告。",
                    "professional_fields": ["项目管理"],
                    "source_material_ids": [],
                },
                source_context=TaskContext(
                    agent_id="agent",
                    session_id="session",
                ),
            ),
        )
        copied = await task_service.create_custom_flow_copy(_OWNER, task.id)
        final_node = copied.nodes[-1].model_copy(
            update={
                "artifact": ArtifactConfig(
                    format=ArtifactFormat.DOCX,
                    filename="示范工程策划稿.docx",
                ),
            },
        )
        custom = await task_service.update_task(
            _OWNER,
            copied.id,
            UpdateTaskRequest(
                knowledge_base_ids=[_KB_ID],
                nodes=[*copied.nodes[:-1], final_node],
            ),
        )
        gateway = TaskKnowledgeGateway(
            _PlanningTemplateKnowledgeBaseFixture(mode="project_materials"),
            SimpleNamespace(),
        )

        run = await task_service.create_run(
            _OWNER,
            custom.id,
            TaskRunRequest(),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)
        self.assertEqual(run.status, RunStatus.AWAITING_INPUT)
        run = await task_service.answer_run(
            _OWNER,
            run.id,
            TaskInputReply(
                request_id=run.pending_input.id,
                answers={
                    "project_name": "示范工程",
                    "project_brief": "正文必须以授权原文为依据。",
                    "source_material_ids": _PROJECT_DOC_ID,
                },
            ),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)
        self.assertEqual(run.status, RunStatus.AWAITING_INPUT)
        self.assertEqual(run.source_material_ids, [_PROJECT_DOC_ID])
        retrieval = run.node_runs[1].result
        project_retrieval = retrieval["project_material_retrieval"]
        self.assertEqual(project_retrieval["status"], "ready")
        self.assertEqual(
            project_retrieval["evidence"][0]["source"],
            _PROJECT_SOURCE,
        )
        self.assertEqual(
            project_retrieval["evidence"][0]["content_sha256"],
            _PROJECT_HASH,
        )
        self.assertEqual(
            {item["evidence_type"] for item in retrieval["evidence"]},
            {"planning_template", "project_material"},
        )

        run = await task_service.answer_run(
            _OWNER,
            run.id,
            TaskInputReply(
                request_id=run.pending_input.id,
                answers={"planning_feedback": "保留已核实来源并标注待核实项。"},
            ),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)

        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        self.assertEqual(planner.calls, 0)
        self.assertEqual(run.node_runs[4].input, run.node_runs[3].output)
        self.assertTrue(
            run.node_runs[-1].metadata["evidence_ledger"]["passed"]
        )
        self.assertTrue(
            run.node_runs[-1].metadata["final_citation_verification"]["passed"]
        )
        self.assertTrue(
            run.node_runs[-1].metadata["final_citation_verification"][
                "verification_scope"
            ]["semantic_support"]
            == "not_automatically_evaluated"
        )
        self.assertIn("### 来源与证据台账", run.final_output)
        self.assertIn(_PROJECT_SOURCE, run.final_output)
        self.assertIn(_PROJECT_DOC_ID, run.final_output)
        self.assertIn(_PROJECT_HASH, run.final_output)
        self.assertIn("语义支持需人工复核", run.final_output)
        self.assertEqual(len(run.artifacts), 1)
        self.assertEqual(len(executor.files), 1)
        await task_service.shutdown()

    async def _run_s02_occurrence_protocol(self, executor, filename):
        planner = _NeverPlanner()
        task_service = TaskService(
            TaskStore(),
            agentscope_executor=executor,
            preview_executor=executor,
            artifact_writer=executor,
            planner=planner,
        )
        template_service = TaskTemplateService(
            TemplateRegistry(),
            task_service,
            compiler=TaskTemplateCompiler(),
        )
        task = await template_service.instantiate(
            _OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version="2.0.0",
                inputs={
                    "project_name": "Occurrence 协议演示项目",
                    "project_brief": "使用本次授权的项目资料与策划模板。",
                    "professional_fields": ["项目管理"],
                    "source_material_ids": [],
                },
                source_context=TaskContext(
                    agent_id="agent",
                    session_id="session",
                ),
            ),
        )
        copied = await task_service.create_custom_flow_copy(_OWNER, task.id)
        final_node = copied.nodes[-1].model_copy(
            update={
                "artifact": ArtifactConfig(
                    format=ArtifactFormat.DOCX,
                    filename=filename,
                ),
            },
        )
        custom = await task_service.update_task(
            _OWNER,
            copied.id,
            UpdateTaskRequest(
                knowledge_base_ids=[_KB_ID],
                nodes=[*copied.nodes[:-1], final_node],
            ),
        )
        gateway = TaskKnowledgeGateway(
            _PlanningTemplateKnowledgeBaseFixture(mode="project_materials"),
            SimpleNamespace(),
        )
        run = await task_service.create_run(
            _OWNER,
            custom.id,
            TaskRunRequest(),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)
        if run.status != RunStatus.AWAITING_INPUT:
            await task_service.shutdown()
            raise AssertionError(
                f"Expected planning feedback input, got {run.status}: {run.error}"
            )
        run = await task_service.answer_run(
            _OWNER,
            run.id,
            TaskInputReply(
                request_id=run.pending_input.id,
                answers={
                    "project_name": "Occurrence 协议演示项目",
                    "project_brief": "使用授权检索资料。",
                    "source_material_ids": _PROJECT_DOC_ID,
                },
            ),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)
        if run.status != RunStatus.AWAITING_INPUT:
            await task_service.shutdown()
            raise AssertionError(
                f"Expected review input, got {run.status}: {run.error}; "
                f"failed_nodes={[(item.name, item.error) for item in run.node_runs if item.error]}"
            )
        run = await task_service.answer_run(
            _OWNER,
            run.id,
            TaskInputReply(
                request_id=run.pending_input.id,
                answers={"planning_feedback": "保留证据锚点。"},
            ),
            gateway,
        )
        await task_service._active_runs[run.id]
        return await task_service.get_run(_OWNER, run.id), planner, task_service

    async def test_new_occurrence_protocol_runs_through_task_service_and_sop(self):
        executor = _OccurrenceProtocolArtifactExecutor(project_materials=True)
        planner = _NeverPlanner()
        task_service = TaskService(
            TaskStore(),
            agentscope_executor=executor,
            preview_executor=executor,
            artifact_writer=executor,
            planner=planner,
        )
        template_service = TaskTemplateService(
            TemplateRegistry(),
            task_service,
            compiler=TaskTemplateCompiler(),
        )
        task = await template_service.instantiate(
            _OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version="2.0.0",
                inputs={
                    "project_name": "Occurrence 协议演示项目",
                    "project_brief": "使用本次授权的项目资料与策划模板。",
                    "professional_fields": ["项目管理"],
                    "source_material_ids": [],
                },
                source_context=TaskContext(
                    agent_id="agent",
                    session_id="session",
                ),
            ),
        )
        copied = await task_service.create_custom_flow_copy(_OWNER, task.id)
        final_node = copied.nodes[-1].model_copy(
            update={
                "artifact": ArtifactConfig(
                    format=ArtifactFormat.DOCX,
                    filename="Occurrence-协议测试.docx",
                ),
            },
        )
        custom = await task_service.update_task(
            _OWNER,
            copied.id,
            UpdateTaskRequest(
                knowledge_base_ids=[_KB_ID],
                nodes=[*copied.nodes[:-1], final_node],
            ),
        )
        knowledge_service = _PlanningTemplateKnowledgeBaseFixture(
            mode="project_materials"
        )
        gateway = TaskKnowledgeGateway(knowledge_service, SimpleNamespace())

        try:
            run = await task_service.create_run(
                _OWNER, custom.id, TaskRunRequest(), gateway
            )
            await task_service._active_runs[run.id]
            run = await task_service.get_run(_OWNER, run.id)
            self.assertEqual(run.status, RunStatus.AWAITING_INPUT)
            run = await task_service.answer_run(
                _OWNER,
                run.id,
                TaskInputReply(
                    request_id=run.pending_input.id,
                    answers={
                        "project_name": "Occurrence 协议演示项目",
                        "project_brief": "使用授权检索资料。",
                        "source_material_ids": _PROJECT_DOC_ID,
                    },
                ),
                gateway,
            )
            await task_service._active_runs[run.id]
            run = await task_service.get_run(_OWNER, run.id)
            self.assertEqual(run.status, RunStatus.AWAITING_INPUT)
            run = await task_service.answer_run(
                _OWNER,
                run.id,
                TaskInputReply(
                    request_id=run.pending_input.id,
                    answers={"planning_feedback": "保留证据锚点。"},
                ),
                gateway,
            )
            await task_service._active_runs[run.id]
            run = await task_service.get_run(_OWNER, run.id)

            self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
            final = run.node_runs[-1]
            self.assertTrue(final.metadata["evidence_ledger"]["passed"])
            self.assertTrue(
                final.metadata["final_citation_verification"]["passed"]
            )
            self.assertEqual(
                final.metadata["final_citation_verification"][
                    "verification_scope"
                ]["semantic_support"],
                "not_automatically_evaluated",
            )
            self.assertEqual(
                final.metadata["final_citation_verification"][
                    "verification_scope"
                ]["citation_id_set"],
                "passed",
            )
            self.assertEqual(
                final.metadata["final_citation_verification"][
                    "verification_scope"
                ]["source_identity_and_quote"],
                "passed",
            )
            self.assertEqual(run.node_runs[1].result["status"], "ready")
            self.assertEqual(len(run.artifacts), 1)
            self.assertTrue(executor.files[run.artifacts[0].path].startswith(b"PK"))
            self.assertEqual(final.output.count("```task-claim-evidence"), 0)
            self.assertEqual(final.output.count("```task-citations"), 1)
            self.assertIn("CL-area:O-area", final.output)
            self.assertIn("CL-rule:O-rule", final.output)
            self.assertIn(
                "UV-reviewer",
                {
                    item.get("unverified_id")
                    for item in final.metadata["evidence_ledger"]["ledger"]
                },
            )
            unverified_statuses = {
                item.get("unverified_id"): item.get("status")
                for item in final.metadata["evidence_ledger"]["ledger"]
                if item.get("unverified_id")
            }
            self.assertEqual(
                unverified_statuses,
                {
                    "UV-reviewer": "待确认",
                    "UV-meter-boundary": "待补充",
                },
            )
            self.assertEqual(
                {item["citation_id"] for item in final.metadata["evidence_ledger"]["ledger"]},
                {"C1", "C2", "—"},
            )
            self.assertEqual(planner.calls, 0)
        finally:
            await task_service.shutdown()

    async def test_incomplete_new_protocol_fails_closed_through_task_service_and_sop(
        self,
    ):
        executor = _IncompleteOccurrenceProtocolArtifactExecutor(
            project_materials=True
        )
        run, planner, task_service = await self._run_s02_occurrence_protocol(
            executor,
            "不完整 Claim 协议.docx",
        )

        try:
            self.assertEqual(run.status, RunStatus.FAILED)
            final = run.node_runs[-1]
            self.assertIn("evidence_ledger", final.metadata, run.error)
            self.assertFalse(final.metadata["evidence_ledger"]["passed"])
            self.assertFalse(
                any(
                    item["error_code"] == "legacy_compatibility_validation"
                    for item in final.metadata["evidence_ledger"][
                        "all_diagnostics"
                    ]
                )
            )
            diagnostic_codes = {
                item["error_code"]
                for item in final.metadata["evidence_ledger"][
                    "all_diagnostics"
                ]
            }
            self.assertIn("incomplete_claim_block", diagnostic_codes)
            self.assertEqual(
                final.metadata["evidence_ledger"]["verification_scope"][
                    "protocol_parse"
                ],
                "failed",
            )
            raw_citation_result = verify_task_citations(
                executor.raw_final_output,
                run.node_runs[1].result["evidence"],
                authorization_subject=_OWNER,
            )
            self.assertTrue(
                raw_citation_result["passed"],
                raw_citation_result["issues"],
            )
            self.assertEqual(
                executor.raw_final_output.count("```task-citations"),
                1,
            )
            self.assertEqual(run.artifacts, [])
            self.assertEqual(executor.files, {})
            self.assertEqual(planner.calls, 0)
        finally:
            await task_service.shutdown()

    async def test_failed_final_ledger_does_not_create_docx(self):
        executor = _ArtifactExecutor(invalid_final=True)
        task_service = TaskService(
            TaskStore(),
            agentscope_executor=executor,
            preview_executor=executor,
            artifact_writer=executor,
            planner=_NeverPlanner(),
        )
        template_service = TaskTemplateService(
            TemplateRegistry(),
            task_service,
            compiler=TaskTemplateCompiler(),
        )
        task = await template_service.instantiate(
            _OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version="2.0.0",
                inputs={
                    "project_name": "测试项目",
                    "project_brief": "形成项目策划报告。",
                    "professional_fields": [],
                },
                source_context=TaskContext(
                    agent_id="agent",
                    session_id="session",
                ),
            ),
        )
        copied = await task_service.create_custom_flow_copy(_OWNER, task.id)
        final_node = copied.nodes[-1].model_copy(
            update={
                "artifact": ArtifactConfig(
                    format=ArtifactFormat.DOCX,
                    filename="无效引用.docx",
                ),
            },
        )
        custom = await task_service.update_task(
            _OWNER,
            copied.id,
            UpdateTaskRequest(
                knowledge_base_ids=[_KB_ID],
                nodes=[*copied.nodes[:-1], final_node],
            ),
        )
        gateway = TaskKnowledgeGateway(
            _PlanningTemplateKnowledgeBaseFixture(),
            SimpleNamespace(),
        )
        run = await task_service.create_run(
            _OWNER,
            custom.id,
            TaskRunRequest(),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)
        run = await task_service.answer_run(
            _OWNER,
            run.id,
            TaskInputReply(
                request_id=run.pending_input.id,
                answers={
                    "project_name": "测试项目",
                    "project_brief": "形成项目策划报告。",
                },
            ),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)
        run = await task_service.answer_run(
            _OWNER,
            run.id,
            TaskInputReply(
                request_id=run.pending_input.id,
                answers={"planning_feedback": "继续整理。"},
            ),
            gateway,
        )
        await task_service._active_runs[run.id]
        run = await task_service.get_run(_OWNER, run.id)

        self.assertEqual(run.status, RunStatus.FAILED)
        self.assertEqual(
            run.node_runs[4].metadata["citation_protocol_status"],
            "verified",
        )
        self.assertTrue(
            run.node_runs[4].metadata["citation_verification"]["passed"]
        )
        self.assertFalse(
            run.node_runs[-1].metadata["evidence_ledger"]["passed"]
        )
        self.assertEqual(
            run.node_runs[-1].metadata["attempts"],
            1,
            f"Run error={run.error}; final node error={run.node_runs[-1].error}",
        )
        self.assertEqual(
            run.node_runs[-1].metadata["authoritative_output_status"][
                "error_code"
            ],
            "TASK_EVIDENCE_LEDGER_FAILED",
        )
        self.assertTrue(
            run.node_runs[-1].output.startswith(
                "> **Task 服务端验收：失败**"
            )
        )
        self.assertEqual(
            executor.agent_calls.count("整理策划输出"),
            1,
        )
        self.assertEqual(run.artifacts, [])
        self.assertEqual(executor.files, {})
        await task_service.shutdown()


async def _empty_search(user_id, knowledge_base_id, query, top_k=5):
    del user_id, knowledge_base_id, query, top_k
    return []


def _citation_document():
    payload = {
        "citations": [
            {
                "citation_id": "C1",
                "knowledge_base_id": _KB_ID,
                "document_id": _DOCUMENT_ID,
                "chunk_index": 0,
                "content_sha256": _CHUNK_HASH,
                "quote": _SOURCE_TEXT,
            }
        ]
    }
    return (
        f"# 项目策划初稿\n\n{_SOURCE_TEXT}[C1]\n\n"
        "```task-citations\n"
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        + "\n```"
    )
