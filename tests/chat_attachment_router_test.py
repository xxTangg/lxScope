# -*- coding: utf-8 -*-
"""Tests for parser-backed chat attachments."""
import io
from unittest import TestCase

from fastapi import FastAPI
from fastapi.testclient import TestClient

from agentscope.app._router import chat_router
from agentscope.app.deps import get_current_user_id, get_knowledge_parsers
from agentscope.message import TextBlock
from agentscope.rag import ParserBase, Section, TextParser, WordParser


HEADERS = {"X-User-ID": "alice"}


class _DocumentParser(ParserBase):
    """Small deterministic parser used to exercise section formatting."""

    supported_media_types = ["application/x-test-document"]

    @classmethod
    def supported_extensions(cls) -> list[str]:
        """Return the fake document extension."""
        return [".tdoc"]

    async def parse(
        self,
        file: bytes | str,
        filename: str,
    ) -> list[Section]:
        """Return two sections, one carrying boundary metadata."""
        del file
        return [
            Section(
                content=TextBlock(text="first section"),
                source=filename,
                metadata={"page": 1},
            ),
            Section(
                content=TextBlock(text="second section"),
                source=filename,
                metadata={},
            ),
        ]


class _EmptyParser(ParserBase):
    """Parser that finds no readable text."""

    supported_media_types = ["application/x-empty-document"]

    @classmethod
    def supported_extensions(cls) -> list[str]:
        """Return the fake empty-document extension."""
        return [".empty"]

    async def parse(
        self,
        file: bytes | str,
        filename: str,
    ) -> list[Section]:
        """Return one structurally valid but empty section."""
        del file
        return [
            Section(
                content=TextBlock(text=""),
                source=filename,
                metadata={},
            ),
        ]


def _make_client(
    parsers: list[ParserBase],
    *,
    max_bytes: int = 1024,
    max_chars: int = 10_000,
) -> TestClient:
    """Build the smallest app needed to exercise the chat router."""
    app = FastAPI()
    app.state.chat_attachment_max_bytes = max_bytes
    app.state.chat_attachment_max_chars = max_chars
    app.include_router(chat_router)
    app.dependency_overrides[get_current_user_id] = lambda: "alice"
    app.dependency_overrides[get_knowledge_parsers] = lambda: parsers
    return TestClient(app)


class ChatAttachmentRouterTest(TestCase):
    """Verify capability discovery, parsing, and request limits."""

    def test_supported_types_are_advertised(self) -> None:
        """Configured text parsers expose media types and extensions."""
        with _make_client([TextParser(), _DocumentParser()]) as client:
            response = client.get(
                "/chat/attachments/supported_content_types",
                headers=HEADERS,
            )

        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertIn("text/plain", body["media_types"])
        self.assertIn("application/x-test-document", body["media_types"])
        self.assertIn(".txt", body["extensions"])
        self.assertIn(".tdoc", body["extensions"])

    def test_document_is_flattened_with_section_metadata(self) -> None:
        """A parsed document becomes one chat-ready text payload."""
        with _make_client([_DocumentParser()]) as client:
            response = client.post(
                "/chat/attachments/parse",
                files={
                    "file": (
                        "report.tdoc",
                        b"source bytes",
                        "application/octet-stream",
                    ),
                },
                headers=HEADERS,
            )

        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["media_type"], "application/x-test-document")
        self.assertEqual(body["section_count"], 2)
        self.assertIn("[page: 1]\nfirst section", body["text"])
        self.assertIn("second section", body["text"])
        self.assertFalse(body["truncated"])

    def test_real_docx_is_extracted(self) -> None:
        """The production Word parser extracts an uploaded DOCX file."""
        from docx import Document

        document = Document()
        document.add_heading("Quarterly report", level=1)
        document.add_paragraph("Revenue increased by 12 percent.")
        payload = io.BytesIO()
        document.save(payload)

        with _make_client(
            [WordParser(include_image=False)],
            max_bytes=1024 * 1024,
        ) as client:
            response = client.post(
                "/chat/attachments/parse",
                files={
                    "file": (
                        "report.docx",
                        payload.getvalue(),
                        (
                            "application/vnd.openxmlformats-officedocument."
                            "wordprocessingml.document"
                        ),
                    ),
                },
                headers=HEADERS,
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("Quarterly report", response.json()["text"])
        self.assertIn("Revenue increased", response.json()["text"])

    def test_text_is_truncated_at_configured_limit(self) -> None:
        """Large extracted text is bounded before it reaches the model."""
        with _make_client([_DocumentParser()], max_chars=8) as client:
            response = client.post(
                "/chat/attachments/parse",
                files={"file": ("report.tdoc", b"x", None)},
                headers=HEADERS,
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()["truncated"])
        self.assertIn("truncated", response.json()["text"])

    def test_unsupported_extension_is_rejected(self) -> None:
        """Files without a configured parser return HTTP 415."""
        with _make_client([TextParser()]) as client:
            response = client.post(
                "/chat/attachments/parse",
                files={"file": ("archive.zip", b"zip", None)},
                headers=HEADERS,
            )

        self.assertEqual(response.status_code, 415, response.text)

    def test_oversized_attachment_is_rejected(self) -> None:
        """The upload byte cap is enforced before parsing."""
        with _make_client([TextParser()], max_bytes=3) as client:
            response = client.post(
                "/chat/attachments/parse",
                files={"file": ("notes.txt", b"four", "text/plain")},
                headers=HEADERS,
            )

        self.assertEqual(response.status_code, 413, response.text)

    def test_empty_document_is_rejected(self) -> None:
        """A parser result without text does not create a blank attachment."""
        with _make_client([_EmptyParser()]) as client:
            response = client.post(
                "/chat/attachments/parse",
                files={"file": ("blank.empty", b"x", None)},
                headers=HEADERS,
            )

        self.assertEqual(response.status_code, 422, response.text)
        self.assertIn("No readable text", response.json()["detail"])
