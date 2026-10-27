"""R7 acceptance for the nine current, versioned Task SOP templates.

Only model, knowledge and Workspace boundaries use explicit synthetic fakes.
Template compilation, TaskService preflight/scheduling, Task-native Tool
adapters, SOP checkpoints, verification handling and artifact generation run
through the production Task module.
"""

from __future__ import annotations

import asyncio
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
from typing import Any
from uuid import uuid4
from zipfile import ZipFile, is_zipfile
from xml.etree import ElementTree

import pytest

from examples.agent_service.task._executor import StepExecutionResult
from examples.agent_service.task._knowledge import TaskKnowledgeError
from examples.agent_service.task._models import (
    NodeRunStatus,
    RunStatus,
    StepType,
    StepVerdict,
    TaskInputReply,
    TaskPreflightRequest,
    TaskRunControlRequest,
    TaskRunRequest,
    UpdateTaskRequest,
)
from examples.agent_service.task._service import (
    TaskService,
    TaskValidationFailure,
    TaskWorkflowConflictError,
)
from examples.agent_service.task._store import TaskStore, TaskStoreConflictError
from examples.agent_service.task._template_models import TemplateInstantiateRequest
from examples.agent_service.task._template_registry import TemplateRegistry
from examples.agent_service.task._template_service import TaskTemplateService


SCENARIO_IDS = (
    "s01_project_init",
    "s02_project_planning",
    "s03_expert_matching",
    "s04_meeting_preparation",
    "s05_meeting_minutes",
    "s06_opinion_processing",
    "s07_report_generation",
    "s08_archive_transfer",
    "s09_expense_settlement",
)
DEMO_INPUT_FILES = (
    "s01-project-init.v2.json",
    "s02-project-planning.v2.json",
    "s03-expert-matching.v3.json",
    "s04-meeting-preparation.v3.json",
    "s05-meeting-minutes.v3.json",
    "s06-opinion-processing.v3.json",
    "s07-report-generation.v3.json",
    "s08-archive-transfer.v3.json",
    "s09-expense-settlement.v3.json",
)
INPUT_ROOT = (
    Path(__file__).resolve().parents[1]
    / "examples"
    / "agent_service"
    / "task"
    / "docs"
    / "p7b-demo-data"
    / "inputs"
)
SOURCE_ID = "r7-synthetic-authorized-document"
KB_ID = "r7-synthetic-authorized-kb"
FAILURE_STEP_KEYS = {
    "s01_project_init": "extract_project_profile",
    "s02_project_planning": "draft_eight_chapter_report",
    "s03_expert_matching": "explain_matching_reasons",
    "s04_meeting_preparation": "prepare_notification",
    "s05_meeting_minutes": "segment_meeting_topics",
    "s06_opinion_processing": "summarize_opinion_dimensions",
    "s07_report_generation": "draft_report_chapters",
    "s08_archive_transfer": "classify_archive_materials",
    "s09_expense_settlement": "prepare_expense_draft",
}


def _inputs_for(template_id: str, template: Any) -> dict[str, Any]:
    path = INPUT_ROOT / DEMO_INPUT_FILES[SCENARIO_IDS.index(template_id)]
    sample = json.loads(path.read_text(encoding="utf-8"))["inputs"]
    inputs = {
        key: value
        for key, value in sample.items()
        if key in template.input_schema.properties
    }
    if "source_material_ids" in template.input_schema.properties:
        inputs["source_material_ids"] = [SOURCE_ID]
    if template_id == "s04_meeting_preparation":
        inputs["meeting_date"] = "2025-11-06 09:00"
    assert set(template.input_schema.required) <= set(inputs)
    return inputs


def _evidence(
    owner: str,
    *,
    document_id: str,
    source: str,
    content: str,
    evidence_type: str,
) -> dict[str, Any]:
    return {
        "knowledge_base_id": KB_ID,
        "knowledge_base_name": "R7 DEMO/SYNTHETIC authorized KB",
        "document_id": document_id,
        "source": source,
        "chunk_index": 0,
        "content_sha256": hashlib.sha256(content.encode("utf-8")).hexdigest(),
        "source_text": content,
        "source_text_truncated": False,
        "source_type": "knowledge_base_chunk",
        "source_status": "ready",
        "authorization_subject": owner,
        "evidence_type": evidence_type,
        "collected_at": datetime.now(timezone.utc).isoformat(),
    }


class _SyntheticAuthorizedKnowledge:
    """A scoped authorized-source fixture; it is not a real KB service."""

    def __init__(self, owner: str) -> None:
        self.owner = owner
        self.retrieval_calls: list[tuple[str, tuple[str, ...]]] = []
        self.fail_runtime_capability: str | None = None

    def _scope(self, user_id: str, ids: list[str]) -> None:
        if user_id != self.owner or ids != [KB_ID]:
            raise TaskKnowledgeError(
                "knowledge_base_unavailable",
                "R7 synthetic KB permission denied.",
            )

    def _record(self, name: str, source_ids: list[str]) -> None:
        self.retrieval_calls.append((name, tuple(source_ids)))
        if self.fail_runtime_capability == name:
            raise TaskKnowledgeError(
                "retrieval_failed",
                "R7 injected synthetic KB retrieval failure.",
            )

    async def inspect_task_scope(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        source_material_ids: list[str],
    ) -> dict[str, Any]:
        self._scope(user_id, knowledge_base_ids)
        if source_material_ids and source_material_ids != [SOURCE_ID]:
            raise TaskKnowledgeError(
                "source_material_not_found", "Unknown R7 synthetic source."
            )
        return {"knowledge_bases": [{"id": KB_ID, "name": "R7 synthetic KB"}]}

    async def retrieve_task_evidence(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        query: str,
        *,
        source_material_ids: list[str],
        top_k: int,
    ) -> dict[str, Any]:
        self._scope(user_id, knowledge_base_ids)
        assert query and top_k == 8
        if source_material_ids and source_material_ids != [SOURCE_ID]:
            raise TaskKnowledgeError(
                "source_material_not_found", "Unknown R7 synthetic source."
            )
        kind = "report.rag_retrieve" if source_material_ids else "project.similar_search"
        if "规则" in query or "办法" in query:
            kind = "project.rule_check"
        self._record(kind, source_material_ids)
        item = _evidence(
            self.owner,
            document_id=SOURCE_ID,
            source="R7 DEMO/SYNTHETIC 项目资料.md",
            content="R7 synthetic project source; all facts require human review.",
            evidence_type="project_material",
        )
        if kind == "project.rule_check":
            item = _evidence(
                self.owner,
                document_id="r7-synthetic-rule",
                source="R7 DEMO/SYNTHETIC 项目规则候选办法.md",
                content="R7 synthetic rule candidate; applicability is unverified.",
                evidence_type="project_rule",
            )
        return {
            "status": "ready",
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "knowledge_bases": [{"id": KB_ID, "name": "R7 synthetic KB"}],
            "evidence": [item],
        }

    async def retrieve_planning_template_evidence(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        *,
        top_k: int,
    ) -> dict[str, Any]:
        self._scope(user_id, knowledge_base_ids)
        assert top_k == 12
        self._record("planning.template_retrieve", [])
        item = _evidence(
            self.owner,
            document_id="r7-synthetic-planning-template",
            source="R7 DEMO/SYNTHETIC 项目策划模板.md",
            content="R7 synthetic planning template; eight chapters, no approval.",
            evidence_type="planning_template",
        )
        return {
            "status": "ready",
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "knowledge_bases": [{"id": KB_ID, "name": "R7 synthetic KB"}],
            "evidence": [item],
            "template_matches": [{"source": item["source"]}],
            "queries": ["R7 synthetic planning template"],
            "citation_protocol": "Only authorized R7 synthetic sources may be cited.",
        }

    async def retrieve_project_material_evidence(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        *,
        project_name: str,
        source_material_ids: list[str],
        top_k: int,
    ) -> dict[str, Any]:
        self._scope(user_id, knowledge_base_ids)
        assert project_name and source_material_ids == [SOURCE_ID] and top_k == 8
        self._record("planning.project_material", source_material_ids)
        return {
            "status": "ready",
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "knowledge_bases": [{"id": KB_ID, "name": "R7 synthetic KB"}],
            "queries": ["R7 synthetic project materials"],
            "evidence": [
                _evidence(
                    self.owner,
                    document_id=SOURCE_ID,
                    source="R7 DEMO/SYNTHETIC 项目委托资料.md",
                    content="R7 synthetic project goals and boundaries await review.",
                    evidence_type="project_material",
                )
            ],
        }

    async def retrieve(
        self, user_id: str, knowledge_base_ids: list[str], query: str
    ) -> str:
        self._scope(user_id, knowledge_base_ids)
        assert query
        return "R7 DEMO/SYNTHETIC authorized context."

    @staticmethod
    def format_task_evidence(result: dict[str, Any]) -> str:
        return json.dumps(result, ensure_ascii=False, default=str)


def _s02_report() -> str:
    sections = (
        ("第一章 项目概况", "项目名称、委托背景和工作边界待核实。"),
        (
            "第二章 资料清单与可信度",
            "R7 DEMO/SYNTHETIC 项目委托资料.md、R7 DEMO/SYNTHETIC 项目策划模板.md；"
            "资料可信度待确认。",
        ),
        ("第三章 能源基线与计量口径", "统计期间和计量口径待补充。"),
        ("第四章 设备台账现状", "设备字段完整性待核实。"),
        ("第五章 分析方法与工作计划", "现场计划待确认。"),
        ("第六章 方案建议", "建议应根据补充资料调整。"),
        ("第七章 预算与进度边界", "预算状态及周期待核实。"),
        ("第八章 风险、待核实事项与交付清单", "责任人待确认。"),
    )
    return "# R7 DEMO/SYNTHETIC 项目策划报告\n\n" + "\n\n".join(
        f"### {heading}\n{body}" for heading, body in sections
    )


class _SyntheticModelAndWorkspace:
    """Deterministic model/verifier and in-memory test artifact writer."""

    def __init__(self, scenario_id: str, inputs: dict[str, Any]) -> None:
        self.scenario_id = scenario_id
        self.inputs = inputs
        self.fail_node_id: str | None = None
        self.reject_node_id: str | None = None
        self.block_node_id: str | None = None
        self.entered = asyncio.Event()
        self.release = asyncio.Event()
        self.calls: list[tuple[str, str]] = []
        self.verifications: list[str] = []
        self.artifacts: dict[str, bytes] = {}

    async def execute_node(
        self, *, node: Any, previous_output: str, context: Any
    ) -> StepExecutionResult:
        assert node.type is StepType.AGENT, "Task-native Tools must use real adapters"
        assert context.run_id and context.task_id and context.user_id
        self.calls.append((node.id, previous_output))
        if node.id == self.block_node_id:
            self.entered.set()
            await self.release.wait()
        if node.id == self.fail_node_id:
            return StepExecutionResult(
                text="R7 injected synthetic Agent failure.",
                status="failed",
                reason="R7 injected synthetic Agent failure.",
            )
        if self.scenario_id == "s02_project_planning":
            text = _s02_report()
        elif self.scenario_id == "s07_report_generation":
            chapters = self.inputs["required_chapters"]
            text = (
                "# R7 DEMO/SYNTHETIC 报告草稿\n\n"
                "授权来源：R7 DEMO/SYNTHETIC 项目资料.md。引用语义待人工复核。\n\n"
            ) + "\n\n".join(
                f"## {title}\n{title}：仅使用 R7 合成资料，待人工核实。"
                for title in chapters
            )
        else:
            native_history = [
                item
                for item in context.history
                if isinstance(item, dict) and item.get("type") == "tool"
            ]
            native_results = "\n\n".join(
                f"### Task 原生适配结果：{item['name']}\n"
                f"{item['output'][:1200]}"
                for item in native_history[-3:]
            )
            text = (
                f"# R7 DEMO/SYNTHETIC {self.scenario_id} / {node.name}\n\n"
                f"本节点输入：{previous_output[:400]}\n\n"
                f"{native_results}\n\n"
                "仅为合成验收草稿；待业务人员复核，不代表外部业务系统已写入。"
            )
        return StepExecutionResult(
            text=text,
            result={"synthetic_model": True, "node_id": node.id},
        )

    async def verify_node(self, *, node: Any, output: str, context: Any) -> StepVerdict:
        assert context.run_id and output
        self.verifications.append(node.id)
        if node.id == self.reject_node_id:
            return StepVerdict(
                passed=False,
                retryable=False,
                message="R7 deterministic verifier rejected synthetic output.",
            )
        return StepVerdict(
            passed=True,
            message="R7 deterministic verifier accepted synthetic fixture.",
        )

    async def preflight_artifact_writer(self, *, context: Any) -> tuple[bool, str]:
        assert context.user_id
        return True, "R7 synthetic Workspace writer available."

    async def write_artifact(
        self, *, path: str, data: bytes, context: Any
    ) -> int:
        assert context.run_id in path and data
        self.artifacts[path] = data
        return len(data)

    async def read_artifact(self, *, path: str, context: Any) -> bytes:
        assert context.run_id in path
        return self.artifacts[path]


async def _prepare(
    scenario_id: str,
) -> tuple[TaskService, Any, Any, Any, _SyntheticModelAndWorkspace, _SyntheticAuthorizedKnowledge, str, dict[str, Any]]:
    registry = TemplateRegistry()
    template = registry.get_current_template(scenario_id)
    assert template is not None
    inputs = _inputs_for(scenario_id, template)
    owner = f"r7-synthetic-{scenario_id}-{uuid4().hex}"
    executor = _SyntheticModelAndWorkspace(scenario_id, inputs)
    gateway = _SyntheticAuthorizedKnowledge(owner)
    service = TaskService(
        TaskStore(),
        preview_executor=executor,
        artifact_writer=executor,
    )
    task = await TaskTemplateService(registry, service).instantiate(
        owner,
        scenario_id,
        TemplateInstantiateRequest(
            template_version=template.template_version,
            inputs=inputs,
        ),
    )
    if scenario_id in {
        "s01_project_init", "s02_project_planning", "s07_report_generation"
    }:
        task = await service.update_task(
            owner,
            task.id,
            UpdateTaskRequest(knowledge_base_ids=[KB_ID]),
        )
        assert task is not None
    assert task.workflow_mode.value == "template_instance"
    assert task.source_template_version == template.template_version
    assert len(task.nodes) == len(template.steps)
    preflight = await service.preflight_task(
        owner, task.id, TaskPreflightRequest(), gateway
    )
    assert preflight is not None and preflight.ready, (
        scenario_id,
        preflight.model_dump(mode="json") if preflight is not None else None,
    )
    return service, registry, task, template, executor, gateway, owner, inputs


def _answer_value(key: str, inputs: dict[str, Any]) -> str:
    value = inputs.get(key, "R7 DEMO/SYNTHETIC 人工复核意见：无补充")
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


async def _drive(
    service: TaskService,
    owner: str,
    run_id: str,
    gateway: _SyntheticAuthorizedKnowledge,
    inputs: dict[str, Any],
) -> Any:
    for _ in range(24):
        active = service._active_runs.get(run_id)
        if active is not None:
            await asyncio.wait_for(active, timeout=30)
        run = await service.get_run(owner, run_id, gateway)
        assert run is not None
        if run.status in {RunStatus.SUCCEEDED, RunStatus.FAILED, RunStatus.CANCELED}:
            return run
        if run.status is RunStatus.AWAITING_INPUT:
            assert run.pending_input is not None
            await service.answer_run(
                owner,
                run.id,
                TaskInputReply(
                    request_id=run.pending_input.id,
                    answers={
                        question.key: _answer_value(question.key, inputs)
                        for question in run.pending_input.questions
                    },
                ),
                gateway,
            )
            continue
        await asyncio.sleep(0)
    raise AssertionError(f"R7 Run {run_id} did not reach a checkpoint")


async def _start_blocked_agent(
    service: TaskService,
    task: Any,
    executor: _SyntheticModelAndWorkspace,
    gateway: _SyntheticAuthorizedKnowledge,
    owner: str,
    inputs: dict[str, Any],
) -> tuple[Any, int]:
    first_agent = next(node for node in task.nodes if node.type is StepType.AGENT)
    executor.block_node_id = first_agent.id
    created = await service.create_run(
        owner,
        task.id,
        TaskRunRequest(input="R7 DEMO/SYNTHETIC in-flight checkpoint"),
        gateway,
    )
    assert created is not None
    for _ in range(8):
        active_worker = service._active_runs.get(created.id)
        assert active_worker is not None
        entered = asyncio.create_task(executor.entered.wait())
        try:
            finished, _ = await asyncio.wait(
                {active_worker, entered},
                timeout=30,
                return_when=asyncio.FIRST_COMPLETED,
            )
            assert finished, "R7 first Agent did not reach its checkpoint"
            if executor.entered.is_set():
                break
        finally:
            entered.cancel()
            await asyncio.gather(entered, return_exceptions=True)
        waiting = await service.get_run(owner, created.id, gateway)
        assert waiting is not None and waiting.status is RunStatus.AWAITING_INPUT
        assert waiting.pending_input is not None
        await service.answer_run(
            owner,
            created.id,
            TaskInputReply(
                request_id=waiting.pending_input.id,
                answers={
                    question.key: _answer_value(question.key, inputs)
                    for question in waiting.pending_input.questions
                },
            ),
            gateway,
        )
    else:
        raise AssertionError("R7 first Agent was not reached")
    active = await service.get_run(owner, created.id, gateway)
    assert active is not None and active.status is RunStatus.RUNNING
    assert active.node_runs[first_agent.order].status is NodeRunStatus.RUNNING
    return active, first_agent.order


def _artifact_text(data: bytes, format_name: str) -> str:
    if format_name == "markdown":
        return data.decode("utf-8")
    assert is_zipfile(BytesIO(data))
    if format_name == "docx":
        with ZipFile(BytesIO(data)) as archive:
            document = ElementTree.fromstring(archive.read("word/document.xml"))
        return "".join(
            item.text or ""
            for item in document.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
        )
    assert format_name == "xlsx"
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        return "\n".join(
            str(value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if value is not None
        )
    finally:
        workbook.close()


def _trace(
    phase: str,
    scenario_id: str,
    run: Any,
    executor: _SyntheticModelAndWorkspace,
) -> None:
    path = os.environ.get("LXSCOPE_TASK_R7_TRACE")
    if not path:
        return
    record = {
        "phase": phase,
        "scenario_id": scenario_id,
        "run_id": run.id,
        "template_version": run.source_template_version,
        "status": run.status.value,
        "control_mode": run.control_mode,
        "run_version": run.run_version,
        "event_sequence": run.event_sequence,
        "snapshot_fingerprint": run.node_snapshot_fingerprint,
        "nodes": [
            {
                "id": item.node_id,
                "name": item.name,
                "status": item.status.value,
                "attempt": item.attempt,
                "attempt_history": [
                    {"attempt": attempt.attempt, "status": attempt.status}
                    for attempt in item.attempt_history
                ],
            }
            for item in run.node_runs
        ],
        "agent_calls": [node_id for node_id, _ in executor.calls],
        "artifacts": [
            {
                "format": item.format.value,
                "bytes": item.size_bytes,
                "sha256": hashlib.sha256(
                    executor.artifacts[item.path]
                ).hexdigest() if item.path in executor.artifacts else None,
            }
            for item in run.artifacts
        ],
    }
    with Path(path).open("a", encoding="utf-8") as stream:
        stream.write(json.dumps(record, ensure_ascii=False) + "\n")


def _assert_business_results(
    scenario_id: str,
    run: Any,
    inputs: dict[str, Any],
    artifact_text: str,
) -> None:
    results = {
        node.capability_key: record.result
        for node, record in zip(run.nodes, run.node_runs)
        if node.capability_key
    }
    assert "R7 DEMO/SYNTHETIC" in artifact_text
    if scenario_id == "s01_project_init":
        assert results["project.similar_search"]["source_scope"] == (
            "authorized_task_knowledge_bases"
        )
        rule = results["project.rule_check"]
        assert rule["status"] == "candidates_for_review"
        assert rule["automated_compliance_conclusion"] is False
        assert rule["evidence"][0]["authorization_subject"] == run.user_id
    elif scenario_id == "s02_project_planning":
        retrieval = results["planning.template_retrieve"]
        assert retrieval["status"] == "ready"
        assert {item["evidence_type"] for item in retrieval["evidence"]} == {
            "planning_template",
            "project_material",
        }
        assert retrieval["project_material_retrieval"]["evidence"][0][
            "document_id"
        ] == SOURCE_ID
        assert run.delivery_status == "draft_unverified"
        assert run.verification_report["retrieval"]["status"] == "ready"
        assert run.verification_report["source_integrity"]["passed"] is True
        for heading in (
            "项目概况",
            "资料清单与可信度",
            "能源基线与计量口径",
            "设备台账现状",
            "分析方法与工作计划",
            "方案建议",
            "预算与进度边界",
            "风险、待核实事项与交付清单",
        ):
            assert artifact_text.count(heading) >= 1
        assert "项目委托资料.md" in artifact_text
    elif scenario_id == "s03_expert_matching":
        ranking = results["expert.user_rank"]
        assert ranking["status"] == "preliminary_ranking"
        assert ranking["official_expert_database"] is False
        assert ranking["official_conflict_check_performed"] is False
        assert ranking["eligible_candidate_count"] == 2
        assert {item["expert_id"] for item in ranking["excluded_candidates"]} == {
            "DEMO-EXPERT-003"
        }
        assert ranking["requires_human_review"] is True
        assert "DEMO-EXPERT-001" in artifact_text
    elif scenario_id == "s04_meeting_preparation":
        cost = results["meeting.cost_calculate"]
        assert cost["official_fee_policy_verified"] is False
        assert Decimal(cost["total"]) == sum(
            (Decimal(line["amount"]) for line in cost["lines"]),
            Decimal("0"),
        )
        assert cost["total"] == "3560.00"
        checked = results["meeting.material_check"]
        assert checked["status"] == "ready_for_ai_review"
        assert checked["notification_sent"] is False
        assert "3560.00" in run.final_output
    elif scenario_id == "s05_meeting_minutes":
        transcript = results["meeting.transcript_import"]
        assert transcript["status"] == "imported_user_transcript"
        assert transcript["transcription_performed"] is False
        assert transcript["transcript_text"] == inputs["transcript_text"]
        assert inputs["meeting_id"] in run.final_output
    elif scenario_id == "s06_opinion_processing":
        classified = results["opinion.classify"]
        assert classified["status"] == "classified_by_user_supplied_rules"
        assert classified["original_count"] == (
            classified["unique_count"] + classified["duplicate_count"]
        )
        assert classified["duplicate_count"] == 1
        assert classified["requires_human_review"] is True
    elif scenario_id == "s07_report_generation":
        retrieval = results["report.rag_retrieve"]
        assert retrieval["status"] == "ready"
        assert retrieval["source_scope"] == "authorized_task_knowledge_bases"
        assert retrieval["evidence"][0]["document_id"] == SOURCE_ID
        assert retrieval["evidence"][0]["authorization_subject"] == run.user_id
        assert "项目资料.md" in artifact_text
        for heading in inputs["required_chapters"]:
            assert heading in artifact_text
    elif scenario_id == "s08_archive_transfer":
        metadata = results["archive.document_metadata"]
        checked = results["archive.completeness_check"]
        assert metadata["file_content_read"] is False
        assert metadata["archive_system_write_performed"] is False
        assert checked["archive_system_write_performed"] is False
        assert checked["incomplete_count"] == 0
        assert checked["items"]
    else:
        assert scenario_id == "s09_expense_settlement"
        imported = results["expense.invoice_import"]
        dedupe = results["expense.invoice_verify_dedupe"]
        checked = results["expense.rule_check"]
        assert imported["ocr_performed"] is False
        assert imported["invoice_authenticity_verified"] is False
        assert dedupe["duplicate_count"] == 0
        assert checked["total"] == "2120.00"
        assert Decimal(checked["total"]) == sum(
            (
                Decimal(item["total"])
                for item in json.loads(inputs["invoice_records_json"])
            ),
            Decimal("0"),
        )
        assert checked["invoice_authenticity_verified"] is False
        assert checked["formal_approval"] is False
        assert checked["erp_submitted"] is False


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
def test_current_scenario_executes_real_taskservice_sop_and_artifact(
    scenario_id: str,
) -> None:
    async def exercise() -> None:
        service, _, task, template, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        try:
            run = await service.create_run(
                owner,
                task.id,
                TaskRunRequest(input="R7 DEMO/SYNTHETIC 九场景验收"),
                gateway,
            )
            assert run is not None and run.engine == "sop"
            final = await _drive(service, owner, run.id, gateway, inputs)
            assert final.status is RunStatus.SUCCEEDED, (
                scenario_id,
                final.error,
                [(item.name, item.status.value, item.error) for item in final.node_runs],
            )
            assert final.id == run.id
            assert final.source_template_version == template.template_version
            assert final.node_snapshot_fingerprint
            assert [item.status for item in final.node_runs] == [
                NodeRunStatus.SUCCEEDED
            ] * len(template.steps)
            assert all(item.attempt == 1 for item in final.node_runs)
            assert len(final.artifacts) == 1
            artifact = final.artifacts[0]
            assert executor.artifacts[artifact.path]
            artifact_text = _artifact_text(
                executor.artifacts[artifact.path], artifact.format.value
            )
            _assert_business_results(scenario_id, final, inputs, artifact_text)
            for node, record in zip(final.nodes, final.node_runs):
                assert record.output
                if node.type is not StepType.INPUT:
                    assert node.id in executor.verifications
                    assert record.metadata["task_ai_verification"]["passed"] is True
                else:
                    assert record.metadata["task_input_validation"]["passed"] is True
                if node.type is StepType.AGENT and node.order > 0:
                    expected_input = final.node_runs[node.order - 1].output
                    assert (node.id, expected_input) in executor.calls
            events = await service.list_events(owner, run.id)
            assert events and events[0].type == "run.created"
            assert events[-1].type == "run.succeeded"
            assert events[-1].sequence == final.event_sequence
            _trace("full_success", scenario_id, final, executor)
        finally:
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
def test_current_scenario_failed_agent_retries_in_same_run_then_steps_and_resumes(
    scenario_id: str,
) -> None:
    async def exercise() -> None:
        service, _, task, template, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        selected_key = FAILURE_STEP_KEYS[scenario_id]
        selected_id = task.template_step_bindings[selected_key]
        selected_order = next(
            node.order for node in task.nodes if node.id == selected_id
        )
        executor.fail_node_id = selected_id
        try:
            created = await service.create_run(
                owner,
                task.id,
                TaskRunRequest(input="R7 DEMO/SYNTHETIC controlled failure"),
                gateway,
            )
            assert created is not None and created.engine == "sop"
            failed = await _drive(service, owner, created.id, gateway, inputs)
            assert failed.id == created.id and failed.status is RunStatus.FAILED
            assert failed.node_runs[selected_order].status is NodeRunStatus.FAILED
            assert all(
                item.status is NodeRunStatus.SUCCEEDED
                for item in failed.node_runs[:selected_order]
            )
            assert all(
                item.status is NodeRunStatus.PENDING
                for item in failed.node_runs[selected_order + 1 :]
            )
            assert not failed.artifacts and not failed.final_output
            _trace("failed_agent", scenario_id, failed, executor)
            preceding = [
                (item.node_id, item.output, item.result, item.attempt)
                for item in failed.node_runs[:selected_order]
            ]
            preceding_calls = {
                node.id: sum(call_id == node.id for call_id, _ in executor.calls)
                for node in task.nodes[:selected_order]
                if node.type is StepType.AGENT
            }
            request = TaskRunControlRequest(
                request_id=f"r7-failed-retry-{scenario_id}",
                expected_run_version=failed.run_version,
            )
            retried = await service.retry_node(
                owner, failed.id, selected_id, request
            )
            assert retried is not None
            assert retried.id == failed.id
            assert retried.status is RunStatus.QUEUED
            assert retried.control_mode == "paused"
            assert retried.node_runs[selected_order].status is NodeRunStatus.PENDING
            assert retried.node_runs[selected_order].attempt_history[-1].status == "failed"
            assert not retried.artifacts and not retried.final_output
            _trace("retry_reset", scenario_id, retried, executor)
            assert [
                (item.node_id, item.output, item.result, item.attempt)
                for item in retried.node_runs[:selected_order]
            ] == preceding
            duplicate = await service.retry_node(
                owner, failed.id, selected_id, request
            )
            assert duplicate is not None and duplicate.run_version == retried.run_version
            with pytest.raises(TaskStoreConflictError):
                await service.retry_node(
                    owner,
                    failed.id,
                    selected_id,
                    TaskRunControlRequest(
                        request_id=f"r7-stale-retry-{scenario_id}",
                        expected_run_version=failed.run_version,
                    ),
                )

            executor.fail_node_id = None
            step_request = TaskRunControlRequest(
                request_id=f"r7-step-{scenario_id}",
                expected_run_version=retried.run_version,
            )
            stepped = await service.step_run(
                owner, retried.id, step_request, gateway
            )
            assert stepped is not None
            active = service._active_runs.get(retried.id)
            assert active is not None
            await asyncio.wait_for(active, timeout=30)
            checkpoint = await service.get_run(owner, retried.id, gateway)
            assert checkpoint is not None
            assert checkpoint.status is RunStatus.QUEUED
            assert checkpoint.control_mode == "paused"
            assert checkpoint.node_runs[selected_order].status is NodeRunStatus.SUCCEEDED
            assert all(
                item.status is NodeRunStatus.PENDING
                for item in checkpoint.node_runs[selected_order + 1 :]
            )
            assert sum(call_id == selected_id for call_id, _ in executor.calls) == 2
            _trace("single_step_checkpoint", scenario_id, checkpoint, executor)
            for node_id, count in preceding_calls.items():
                assert sum(call_id == node_id for call_id, _ in executor.calls) == count
            duplicate_step = await service.step_run(
                owner, retried.id, step_request, gateway
            )
            assert duplicate_step is not None
            assert duplicate_step.run_version == checkpoint.run_version

            store = service._store
            await service.shutdown()
            service = TaskService(
                store,
                preview_executor=executor,
                artifact_writer=executor,
            )
            parked = await service.get_run(owner, created.id, gateway)
            assert parked is not None and parked.control_mode == "paused"
            assert created.id not in service._active_runs
            continued = await service.continue_run(
                owner,
                created.id,
                TaskRunControlRequest(
                    request_id=f"r7-continue-{scenario_id}",
                    expected_run_version=parked.run_version,
                ),
                gateway,
            )
            assert continued is not None
            final = await _drive(service, owner, created.id, gateway, inputs)
            assert final.status is RunStatus.SUCCEEDED, (
                scenario_id,
                final.error,
                [(item.status.value, item.error) for item in final.node_runs],
            )
            assert final.id == failed.id
            assert final.source_template_version == template.template_version
            assert [
                (item.node_id, item.output, item.result, item.attempt)
                for item in final.node_runs[:selected_order]
            ] == preceding
            selected = final.node_runs[selected_order]
            assert selected.attempt == 2
            assert [item.status for item in selected.attempt_history] == [
                "failed",
                "succeeded",
            ]
            assert all(
                item.status is NodeRunStatus.SUCCEEDED and item.attempt == 1
                for item in final.node_runs[selected_order + 1 :]
            )
            assert len(final.artifacts) == 1
            assert executor.artifacts[final.artifacts[0].path]
            events = await service.list_events(owner, created.id)
            assert any(event.type == "node.retry_requested" for event in events)
            assert any(event.type == "run.step_completed" for event in events)
            assert any(event.type == "run.continued" for event in events)
            assert events[-1].type == "run.succeeded"
            assert events[-1].sequence == final.event_sequence
            _trace("retry_completed", scenario_id, final, executor)
        finally:
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
def test_current_scenario_pauses_in_flight_and_continues_after_service_rebuild(
    scenario_id: str,
) -> None:
    async def exercise() -> None:
        service, _, task, template, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        try:
            active, agent_order = await _start_blocked_agent(
                service, task, executor, gateway, owner, inputs
            )
            request = TaskRunControlRequest(
                request_id=f"r7-inflight-pause-{scenario_id}",
                expected_run_version=active.run_version,
            )
            requested = await service.pause_run(owner, active.id, request)
            assert requested is not None and requested.control_mode == "pause_requested"
            executor.release.set()
            worker = service._active_runs.get(active.id)
            assert worker is not None
            await asyncio.wait_for(worker, timeout=30)
            paused = await service.get_run(owner, active.id, gateway)
            assert paused is not None
            assert paused.status is RunStatus.QUEUED
            assert paused.control_mode == "paused"
            assert all(
                item.status is NodeRunStatus.SUCCEEDED
                for item in paused.node_runs[: agent_order + 1]
            )
            assert all(
                item.status is NodeRunStatus.PENDING
                for item in paused.node_runs[agent_order + 1 :]
            )
            assert not paused.artifacts
            _trace("inflight_paused", scenario_id, paused, executor)
            prefix = [
                (item.node_id, item.output, item.attempt)
                for item in paused.node_runs[: agent_order + 1]
            ]
            agent_calls = len(executor.calls)
            duplicate = await service.pause_run(owner, active.id, request)
            assert duplicate is not None and duplicate.run_version == paused.run_version
            store = service._store
            await service.shutdown()
            service = TaskService(
                store, preview_executor=executor, artifact_writer=executor
            )
            parked = await service.get_run(owner, active.id, gateway)
            assert parked is not None and parked.control_mode == "paused"
            assert active.id not in service._active_runs
            assert len(executor.calls) == agent_calls
            resumed = await service.continue_run(
                owner,
                active.id,
                TaskRunControlRequest(
                    request_id=f"r7-inflight-continue-{scenario_id}",
                    expected_run_version=parked.run_version,
                ),
                gateway,
            )
            assert resumed is not None
            final = await _drive(service, owner, active.id, gateway, inputs)
            assert final.status is RunStatus.SUCCEEDED
            assert final.source_template_version == template.template_version
            assert [
                (item.node_id, item.output, item.attempt)
                for item in final.node_runs[: agent_order + 1]
            ] == prefix
            assert sum(
                node_id == task.nodes[agent_order].id
                for node_id, _ in executor.calls
            ) == 1
            events = await service.list_events(owner, active.id)
            assert any(event.type == "run.pause_requested" for event in events)
            assert any(event.type == "run.paused" for event in events)
            assert any(event.type == "run.continued" for event in events)
            assert events[-1].sequence == final.event_sequence
            _trace("inflight_resumed", scenario_id, final, executor)
        finally:
            executor.release.set()
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
def test_current_scenario_stop_in_flight_is_terminal(scenario_id: str) -> None:
    async def exercise() -> None:
        service, _, task, _, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        try:
            active, agent_order = await _start_blocked_agent(
                service, task, executor, gateway, owner, inputs
            )
            request = TaskRunControlRequest(
                request_id=f"r7-stop-{scenario_id}",
                expected_run_version=active.run_version,
            )
            stopped = await service.stop_run(owner, active.id, request)
            assert stopped is not None
            worker = service._active_runs.get(active.id)
            assert worker is not None
            await asyncio.wait_for(worker, timeout=30)
            final = await service.get_run(owner, active.id, gateway)
            assert final is not None and final.status is RunStatus.CANCELED
            assert final.control_mode == "stopped"
            assert all(
                item.status is NodeRunStatus.CANCELED and item.attempt == 0
                for item in final.node_runs[agent_order + 1 :]
            )
            assert not final.artifacts
            _trace("inflight_stopped", scenario_id, final, executor)
            duplicate = await service.stop_run(owner, active.id, request)
            assert duplicate is not None and duplicate.run_version == final.run_version
            with pytest.raises(TaskWorkflowConflictError):
                await service.continue_run(
                    owner,
                    active.id,
                    TaskRunControlRequest(
                        request_id=f"r7-continue-stopped-{scenario_id}",
                        expected_run_version=final.run_version,
                    ),
                    gateway,
                )
            events = await service.list_events(owner, active.id)
            assert any(event.type in {"run.canceled", "run.stopped"} for event in events)
        finally:
            executor.release.set()
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
def test_current_scenario_input_wait_survives_service_rebuild(
    scenario_id: str,
) -> None:
    async def exercise() -> None:
        service, _, task, template, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        try:
            created = await service.create_run(
                owner,
                task.id,
                TaskRunRequest(input="R7 DEMO/SYNTHETIC persisted human input"),
                gateway,
            )
            assert created is not None
            worker = service._active_runs.get(created.id)
            assert worker is not None
            await asyncio.wait_for(worker, timeout=30)
            waiting = await service.get_run(owner, created.id, gateway)
            assert waiting is not None
            assert waiting.status is RunStatus.AWAITING_INPUT
            assert waiting.pending_input is not None
            pending_id = waiting.pending_input.id
            assert waiting.node_runs[0].status is NodeRunStatus.AWAITING_INPUT
            assert all(item.attempt == 0 for item in waiting.node_runs[1:])
            _trace("input_waiting", scenario_id, waiting, executor)
            store = service._store
            await service.shutdown()
            service = TaskService(
                store, preview_executor=executor, artifact_writer=executor
            )
            restored = await service.get_run(owner, created.id, gateway)
            assert restored is not None
            assert restored.status is RunStatus.AWAITING_INPUT
            assert restored.pending_input is not None
            assert restored.pending_input.id == pending_id
            assert created.id not in service._active_runs
            answered = await service.answer_run(
                owner,
                created.id,
                TaskInputReply(
                    request_id=pending_id,
                    answers={
                        question.key: _answer_value(question.key, inputs)
                        for question in restored.pending_input.questions
                    },
                ),
                gateway,
            )
            assert answered is not None
            final = await _drive(service, owner, created.id, gateway, inputs)
            assert final.status is RunStatus.SUCCEEDED
            assert final.id == created.id
            assert final.source_template_version == template.template_version
            assert final.node_runs[0].attempt == 1
            assert final.node_runs[0].metadata["task_input_validation"]["passed"] is True
            events = await service.list_events(owner, created.id)
            assert any(event.type == "run.awaiting_input" for event in events)
            assert any(event.type == "input.answered" for event in events)
            _trace("input_restored", scenario_id, final, executor)
        finally:
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
def test_current_scenario_rejects_failed_ai_verification_without_artifact(
    scenario_id: str,
) -> None:
    async def exercise() -> None:
        service, _, task, _, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        executor.reject_node_id = task.nodes[-1].id
        try:
            created = await service.create_run(
                owner,
                task.id,
                TaskRunRequest(input="R7 DEMO/SYNTHETIC verifier rejection"),
                gateway,
            )
            assert created is not None
            failed = await _drive(service, owner, created.id, gateway, inputs)
            assert failed.status is RunStatus.FAILED
            assert failed.node_runs[-1].status is NodeRunStatus.FAILED
            assert failed.node_runs[-1].metadata["task_ai_verification"]["passed"] is False
            assert not failed.artifacts and not executor.artifacts
            assert not failed.final_output
            _trace("ai_rejected", scenario_id, failed, executor)
        finally:
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize(
    "scenario_id",
    ("s01_project_init", "s02_project_planning", "s07_report_generation"),
)
def test_knowledge_scenarios_reject_unauthorized_scope_before_run(
    scenario_id: str,
) -> None:
    async def exercise() -> None:
        service, _, task, _, _, gateway, owner, _ = await _prepare(scenario_id)
        gateway.owner = "r7-synthetic-unrelated-owner"
        try:
            preflight = await service.preflight_task(
                owner, task.id, TaskPreflightRequest(), gateway
            )
            assert preflight is not None and preflight.ready is False
            with pytest.raises(TaskValidationFailure):
                await service.create_run(
                    owner,
                    task.id,
                    TaskRunRequest(input="R7 unauthorized knowledge scope"),
                    gateway,
                )
            assert await service.list_runs(owner, task.id) == []
        finally:
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
def test_current_scenario_tool_retry_requires_confirmation_and_invalidates_suffix(
    scenario_id: str,
) -> None:
    async def exercise() -> None:
        service, _, task, _, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        tool = next(node for node in task.nodes if node.type is StepType.TOOL)
        executor.fail_node_id = task.nodes[-1].id
        try:
            created = await service.create_run(
                owner,
                task.id,
                TaskRunRequest(input="R7 DEMO/SYNTHETIC downstream invalidation"),
                gateway,
            )
            assert created is not None
            failed = await _drive(service, owner, created.id, gateway, inputs)
            assert failed.status is RunStatus.FAILED
            assert failed.node_runs[tool.order].status is NodeRunStatus.SUCCEEDED
            assert failed.node_runs[-1].status is NodeRunStatus.FAILED
            assert all(
                item.status is NodeRunStatus.SUCCEEDED
                for item in failed.node_runs[:tool.order]
            )
            prefix = [
                (item.node_id, item.output, item.attempt)
                for item in failed.node_runs[:tool.order]
            ]
            request_id = f"r7-confirmed-tool-retry-{scenario_id}"
            with pytest.raises(TaskWorkflowConflictError):
                await service.retry_node(
                    owner,
                    created.id,
                    tool.id,
                    TaskRunControlRequest(
                        request_id=request_id,
                        expected_run_version=failed.run_version,
                    ),
                )
            confirmed_request = TaskRunControlRequest(
                request_id=request_id,
                expected_run_version=failed.run_version,
                side_effects_verified=True,
            )
            reset = await service.retry_node(
                owner, created.id, tool.id, confirmed_request
            )
            assert reset is not None and reset.status is RunStatus.QUEUED
            assert reset.control_mode == "paused"
            assert [
                (item.node_id, item.output, item.attempt)
                for item in reset.node_runs[:tool.order]
            ] == prefix
            assert all(
                item.status is NodeRunStatus.PENDING and not item.output
                for item in reset.node_runs[tool.order:]
            )
            assert reset.node_runs[tool.order].attempt_history[-1].status == "succeeded"
            assert reset.node_runs[-1].attempt_history[-1].status == "failed"
            assert not reset.artifacts and not reset.final_output
            _trace("tool_suffix_invalidated", scenario_id, reset, executor)
            duplicate = await service.retry_node(
                owner, created.id, tool.id, confirmed_request
            )
            assert duplicate is not None and duplicate.run_version == reset.run_version
            executor.fail_node_id = None
            await service.continue_run(
                owner,
                created.id,
                TaskRunControlRequest(
                    request_id=f"r7-continue-tool-{scenario_id}",
                    expected_run_version=reset.run_version,
                ),
                gateway,
            )
            final = await _drive(service, owner, created.id, gateway, inputs)
            assert final.status is RunStatus.SUCCEEDED
            assert final.id == failed.id
            assert [
                (item.node_id, item.output, item.attempt)
                for item in final.node_runs[:tool.order]
            ] == prefix
            assert final.node_runs[tool.order].attempt == 2
            assert final.node_runs[tool.order].metadata["side_effects_verified"]["confirmed_by"] == owner
            assert final.node_runs[-1].attempt == 2
            assert len(final.artifacts) == 1
            events = await service.list_events(owner, created.id)
            retry_event = next(
                item for item in events if item.type == "node.retry_requested"
            )
            assert set(retry_event.payload["invalidated_node_ids"]) == {
                item.id for item in task.nodes[tool.order + 1 :]
            }
            assert retry_event.payload["side_effects_verified"] is True
            _trace("tool_retry_completed", scenario_id, final, executor)
        finally:
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
@pytest.mark.parametrize("corruption", ("fingerprint", "node_identity"))
def test_current_scenario_corrupt_checkpoint_fails_closed_after_real_prefix(
    scenario_id: str,
    corruption: str,
) -> None:
    async def exercise() -> None:
        service, _, task, _, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        try:
            active, agent_order = await _start_blocked_agent(
                service, task, executor, gateway, owner, inputs
            )
            await service.pause_run(
                owner,
                active.id,
                TaskRunControlRequest(
                    request_id=f"r7-corrupt-pause-{corruption}-{scenario_id}",
                    expected_run_version=active.run_version,
                ),
            )
            executor.release.set()
            worker = service._active_runs.get(active.id)
            assert worker is not None
            await asyncio.wait_for(worker, timeout=30)
            checkpoint = await service.get_run(owner, active.id, gateway)
            assert checkpoint is not None and checkpoint.control_mode == "paused"
            assert checkpoint.node_runs[agent_order].status is NodeRunStatus.SUCCEEDED
            assert all(
                item.status is NodeRunStatus.PENDING
                for item in checkpoint.node_runs[agent_order + 1 :]
            )
            if corruption == "fingerprint":
                patch = {"node_snapshot_fingerprint": "f" * 64}
            else:
                state = dict(checkpoint.sop_state)
                identities = list(state["task_node_ids"])
                identities[0] = "r7-synthetic-corrupt-node-id"
                state["task_node_ids"] = identities
                patch = {"sop_state": state}
            damaged = await service._store.transition_run(
                owner,
                active.id,
                expected_version=checkpoint.run_version,
                run_patch=patch,
            )
            store = service._store
            await service.shutdown()
            service = TaskService(
                store, preview_executor=executor, artifact_writer=executor
            )
            await service.continue_run(
                owner,
                active.id,
                TaskRunControlRequest(
                    request_id=f"r7-corrupt-continue-{corruption}-{scenario_id}",
                    expected_run_version=damaged.run_version,
                ),
                gateway,
            )
            restarted = service._active_runs.get(active.id)
            assert restarted is not None
            await asyncio.wait_for(restarted, timeout=30)
            failed = await service.get_run(owner, active.id, gateway)
            assert failed is not None and failed.status is RunStatus.FAILED
            assert all(
                item.status is NodeRunStatus.PENDING and item.attempt == 0
                for item in failed.node_runs[agent_order + 1 :]
            )
            assert not failed.artifacts
            assert sum(
                node_id == task.nodes[agent_order].id
                for node_id, _ in executor.calls
            ) == 1
            _trace(f"corrupt_{corruption}_rejected", scenario_id, failed, executor)
        finally:
            executor.release.set()
            await service.shutdown()

    asyncio.run(exercise())


@pytest.mark.parametrize("scenario_id", SCENARIO_IDS)
def test_current_scenario_frozen_run_survives_instance_version_drift(
    scenario_id: str,
) -> None:
    async def exercise() -> None:
        service, _, task, template, executor, gateway, owner, inputs = (
            await _prepare(scenario_id)
        )
        try:
            active, agent_order = await _start_blocked_agent(
                service, task, executor, gateway, owner, inputs
            )
            await service.pause_run(
                owner,
                active.id,
                TaskRunControlRequest(
                    request_id=f"r7-version-pause-{scenario_id}",
                    expected_run_version=active.run_version,
                ),
            )
            executor.release.set()
            worker = service._active_runs.get(active.id)
            assert worker is not None
            await asyncio.wait_for(worker, timeout=30)
            checkpoint = await service.get_run(owner, active.id, gateway)
            assert checkpoint is not None and checkpoint.control_mode == "paused"
            assert checkpoint.source_template_version == template.template_version
            fingerprint = checkpoint.node_snapshot_fingerprint
            assert fingerprint
            current_task = await service._store.get_task(owner, task.id)
            assert current_task is not None
            # Revise only this test-owned Task instance, never the registry template.
            await service._store.save_task(
                current_task.model_copy(
                    update={
                        "revision": current_task.revision + 1,
                        "source_template_version": "99.0.0-r7-test-instance",
                    }
                )
            )
            store = service._store
            await service.shutdown()
            service = TaskService(
                store, preview_executor=executor, artifact_writer=executor
            )
            parked = await service.get_run(owner, active.id, gateway)
            assert parked is not None
            assert parked.source_template_version == template.template_version
            assert parked.node_snapshot_fingerprint == fingerprint
            await service.continue_run(
                owner,
                active.id,
                TaskRunControlRequest(
                    request_id=f"r7-version-continue-{scenario_id}",
                    expected_run_version=parked.run_version,
                ),
                gateway,
            )
            final = await _drive(service, owner, active.id, gateway, inputs)
            assert final.status is RunStatus.SUCCEEDED
            assert final.id == checkpoint.id
            assert final.source_template_version == template.template_version
            assert final.node_snapshot_fingerprint == fingerprint
            assert final.node_runs[agent_order].attempt == 1
            _trace("frozen_version_resumed", scenario_id, final, executor)
        finally:
            executor.release.set()
            await service.shutdown()

    asyncio.run(exercise())
