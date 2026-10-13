"""S02 v5 lightweight draft flow using the existing Task SOP and writer."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
from datetime import datetime, timezone
from typing import Any

from docx import Document
import pytest

from examples.agent_service.task._executor import StepExecutionResult
from examples.agent_service.task._models import (
    RunStatus,
    StepVerdict,
    TaskContext,
    TaskInputReply,
    TaskRunRequest,
    TaskToolSchema,
    UpdateTaskRequest,
)
from examples.agent_service.task._s02_lightweight_output import (
    validate_s02_lightweight_markdown,
)
from examples.agent_service.task._service import TaskService
from examples.agent_service.task._store import TaskStore
from examples.agent_service.task._template_compiler import TaskTemplateCompiler
from examples.agent_service.task._template_models import (
    TemplateInstantiateRequest,
)
from examples.agent_service.task._template_registry import TemplateRegistry
from examples.agent_service.task._template_service import TaskTemplateService


OWNER = "demo-s02-v5-owner"
OTHER_OWNER = "demo-s02-v5-other-owner"
KB_ID = "demo-s02-v5-kb"
PROJECT_DOCUMENT_ID = "demo-s02-v5-project"
TEMPLATE_DOCUMENT_ID = "demo-s02-v5-template"


def _evidence(
    *,
    evidence_type: str,
    document_id: str,
    source: str,
    text: str,
    authorization_subject: str = OWNER,
    knowledge_base_id: str = KB_ID,
    source_status: str = "ready",
) -> dict[str, Any]:
    return {
        "knowledge_base_id": knowledge_base_id,
        "knowledge_base_name": "DEMO/SYNTHETIC S02 v5 test KB",
        "document_id": document_id,
        "source": source,
        "chunk_index": 0,
        "content_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "source_text": text,
        "source_text_truncated": False,
        "source_type": "knowledge_base_chunk",
        "source_status": source_status,
        "authorization_subject": authorization_subject,
        "evidence_type": evidence_type,
        "collected_at": datetime.now(timezone.utc).isoformat(),
    }


def _snapshot() -> list[dict[str, Any]]:
    return [
        _evidence(
            evidence_type="project_material",
            document_id=PROJECT_DOCUMENT_ID,
            source="DEMO/SYNTHETIC 项目委托资料.md",
            text=(
                "DEMO/SYNTHETIC 项目委托资料。项目目标、范围和边界参数"
                "均为演示输入，未完成现场核实。"
            ),
        ),
        _evidence(
            evidence_type="planning_template",
            document_id=TEMPLATE_DOCUMENT_ID,
            source="DEMO/SYNTHETIC 项目策划模板.md",
            text=(
                "DEMO/SYNTHETIC 项目策划模板，供报告结构参考；"
                "不代表正式制度、审批或监管要求。"
            ),
        ),
    ]


def _report() -> str:
    sections = (
        ("第一章 项目概况", "项目名称、委托背景和工作边界待核实。"),
        ("第二章 资料清单与可信度", "资料来源及其可信度待确认。"),
        ("第三章 能源基线与计量口径", "统计期间、计量口径和表计边界待补充。"),
        ("第四章 设备台账现状", "设备字段完整性和现场情况待核实。"),
        ("第五章 分析方法与工作计划", "分析假设、访谈安排和现场计划待确认。"),
        ("第六章 方案建议", "后续建议应根据补充资料调整。"),
        ("第七章 预算与进度边界", "预算状态及计划周期待核实。"),
        ("第八章 风险、待核实事项与交付清单", "风险责任人与复核人待确认。"),
    )
    return "# DEMO/SYNTHETIC 项目策划报告\n\n" + "\n\n".join(
        f"### {title}\n{body}" for title, body in sections
    )


class _NoPlanner:
    async def plan(self, **_: Any) -> Any:
        raise AssertionError("S02 template instantiation must not call Planner")


class _DemoKnowledgeGateway:
    """Synthetic authorized-source fixture, not a live RAG integration."""

    def __init__(self, evidence: list[dict[str, Any]] | None = None) -> None:
        self.evidence = evidence or _snapshot()

    @property
    def template_evidence(self) -> list[dict[str, Any]]:
        return [item for item in self.evidence if item["evidence_type"] == "planning_template"]

    @property
    def project_evidence(self) -> list[dict[str, Any]]:
        return [item for item in self.evidence if item["evidence_type"] == "project_material"]

    async def inspect_task_scope(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        source_material_ids: list[str],
    ) -> dict[str, Any]:
        assert user_id == OWNER
        assert knowledge_base_ids == [KB_ID]
        assert source_material_ids in ([], [PROJECT_DOCUMENT_ID])
        return {"knowledge_bases": [{"id": KB_ID, "name": "DEMO/SYNTHETIC S02 v5 test KB"}]}

    async def retrieve_planning_template_evidence(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        *,
        top_k: int,
    ) -> dict[str, Any]:
        assert user_id == OWNER and knowledge_base_ids == [KB_ID] and top_k == 12
        return {
            "status": "ready",
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "knowledge_bases": [{"id": KB_ID, "name": "DEMO/SYNTHETIC S02 v5 test KB"}],
            "evidence": self.template_evidence,
            "template_matches": [{"source": item["source"]} for item in self.template_evidence],
            "queries": ["DEMO/SYNTHETIC planning template"],
            "citation_protocol": "Evidence snapshot is server-derived.",
        }

    async def retrieve_project_material_evidence(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        *,
        project_name: str,
        source_material_ids: list[str],
        top_k: int,
    ) -> dict[str, Any]:
        assert user_id == OWNER and knowledge_base_ids == [KB_ID]
        assert project_name and source_material_ids == [PROJECT_DOCUMENT_ID]
        assert top_k == 8
        return {
            "status": "ready",
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "knowledge_bases": [{"id": KB_ID, "name": "DEMO/SYNTHETIC S02 v5 test KB"}],
            "evidence": self.project_evidence,
            "queries": ["DEMO/SYNTHETIC project materials"],
        }

    @staticmethod
    def format_task_evidence(result: dict[str, Any]) -> str:
        return json.dumps(result, ensure_ascii=False, default=str)


class _DemoS02Executor:
    def __init__(self) -> None:
        self.artifact_writes: dict[str, bytes] = {}
        self.agent_calls: list[str] = []

    async def list_tools(self, _context: TaskContext) -> list[TaskToolSchema]:
        return []

    async def execute_node(
        self,
        *,
        node: Any,
        previous_output: str,
        context: Any,
    ) -> StepExecutionResult:
        del previous_output, context
        self.agent_calls.append(node.name)
        return StepExecutionResult(
            text=_report(),
            result={"status": "completed", "output": _report()},
        )

    async def verify_node(self, **_: Any) -> StepVerdict:
        return StepVerdict(passed=True)

    async def preflight_artifact_writer(self, *, context: Any) -> tuple[bool, str]:
        assert context.user_id == OWNER
        return True, "DEMO/SYNTHETIC Workspace fixture available"

    async def write_artifact(
        self,
        *,
        path: str,
        data: bytes,
        context: Any,
    ) -> int:
        assert context.user_id == OWNER
        self.artifact_writes[path] = data
        return len(data)

    async def read_artifact(self, *, path: str, context: Any) -> bytes:
        assert context.user_id == OWNER
        return self.artifact_writes[path]


def _exercise_s02_v5(
    evidence: list[dict[str, Any]] | None = None,
) -> tuple[Any, _DemoS02Executor, Any, Any]:
    async def exercise() -> tuple[Any, _DemoS02Executor, Any, Any]:
        executor = _DemoS02Executor()
        gateway = _DemoKnowledgeGateway(evidence)
        service = TaskService(
            TaskStore(),
            agentscope_executor=executor,  # type: ignore[arg-type]
            preview_executor=executor,  # type: ignore[arg-type]
            planner=_NoPlanner(),  # type: ignore[arg-type]
            artifact_writer=executor,  # type: ignore[arg-type]
        )
        try:
            task = await TaskTemplateService(
                TemplateRegistry(), service, compiler=TaskTemplateCompiler()
            ).instantiate(
                OWNER,
                "s02_project_planning",
                TemplateInstantiateRequest(
                    template_version="5.0.0",
                    inputs={
                        "project_name": "DEMO/SYNTHETIC 青岚园区策划演示项目",
                        "project_brief": "用于验证轻量草稿 Task 链路。",
                        "source_material_ids": [PROJECT_DOCUMENT_ID],
                    },
                    source_context=TaskContext(
                        agent_id="demo-s02-v5-agent",
                        session_id="demo-s02-v5-session",
                    ),
                ),
            )
            task = await service.update_task(
                OWNER,
                task.id,
                UpdateTaskRequest(knowledge_base_ids=[KB_ID]),
            )
            run = await service.create_run(
                OWNER,
                task.id,
                TaskRunRequest(),
                gateway,  # type: ignore[arg-type]
            )
            assert run is not None
            for _ in range(10):
                active = service._active_runs.get(run.id)
                if active is not None:
                    await active
                run = await service.get_run(OWNER, run.id)
                assert run is not None
                if run.status in {RunStatus.SUCCEEDED, RunStatus.FAILED}:
                    artifact_read = None
                    other_user_read = None
                    if run.artifacts:
                        artifact_read = await service.read_artifact(
                            OWNER,
                            run.id,
                            run.artifacts[0].id,
                        )
                        other_user_read = await service.read_artifact(
                            OTHER_OWNER,
                            run.id,
                            run.artifacts[0].id,
                        )
                    return run, executor, artifact_read, other_user_read
                assert run.status is RunStatus.AWAITING_INPUT
                assert run.pending_input is not None
                answers: dict[str, str] = {}
                for question in run.pending_input.questions:
                    if question.key == "project_name":
                        answers[question.key] = "DEMO/SYNTHETIC 青岚园区策划演示项目"
                    elif question.key == "project_brief":
                        answers[question.key] = "用于验证轻量草稿 Task 链路。"
                    elif question.key == "source_material_ids":
                        answers[question.key] = json.dumps([PROJECT_DOCUMENT_ID])
                    elif question.key == "planning_feedback":
                        answers[question.key] = "无补充"
                    elif question.required:
                        answers[question.key] = "DEMO/SYNTHETIC test input"
                await service.answer_run(
                    OWNER,
                    run.id,
                    TaskInputReply(
                        request_id=run.pending_input.id,
                        answers=answers,
                    ),
                    gateway,  # type: ignore[arg-type]
                )
            raise AssertionError("S02 v5 did not reach a terminal state")
        finally:
            await service.shutdown()

    return asyncio.run(exercise())


def test_s02_v5_lightweight_markdown_shape_checks_eight_chapters() -> None:
    assert validate_s02_lightweight_markdown(_report())["passed"] is True
    missing_chapter = _report().replace(
        "### 第六章 方案建议\n后续建议应根据补充资料调整。\n\n",
        "",
    )
    result = validate_s02_lightweight_markdown(missing_chapter)
    assert result["passed"] is False
    assert result["error_code"] == "TASK_S02_LIGHTWEIGHT_OUTPUT_INVALID"


def test_s02_v5_taskservice_sop_generates_run_linked_draft_docx() -> None:
    run, executor, artifact_read, other_user_read = _exercise_s02_v5()

    assert run.status is RunStatus.SUCCEEDED
    assert run.source_template_id == "s02_project_planning"
    assert run.source_template_version == "5.0.0"
    assert run.delivery_mode == "lightweight_draft"
    assert run.delivery_status == "draft_unverified"
    assert run.verification_report["delivery_profile"] == "s02_lightweight_draft"
    assert run.verification_report["evidence_ledger_status"] == "not_performed"
    assert "evidence_ledger_passed" not in run.verification_report
    assert run.verification_report["chapter_count"] == 8
    assert run.verification_report["source_integrity"]["passed"] is True
    assert run.verification_report["retrieval"]["status"] == "ready"
    assert run.verification_report["retrieval"]["matched_source_count"] == 2
    assert len(run.artifacts) == 1
    artifact = run.artifacts[0]
    assert artifact.run_id == run.id
    assert artifact.delivery_status == "draft_unverified"
    assert artifact_read is not None
    assert artifact_read[0].id == artifact.id
    assert artifact_read[1] == next(iter(executor.artifact_writes.values()))
    assert other_user_read is None
    assert len(executor.agent_calls) == 2
    assert len(executor.artifact_writes) == 1
    docx_bytes = next(iter(executor.artifact_writes.values()))
    assert docx_bytes.startswith(b"PK")
    document = Document(io.BytesIO(docx_bytes))
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    label = "项目策划初稿 · 尚未完成正式核验"
    assert label in all_text
    assert any(
        label in paragraph.text
        for section in document.sections
        for paragraph in section.header.paragraphs
    )
    assert "task-claim-evidence" not in run.final_output
    assert "task-citations" not in run.final_output


@pytest.mark.parametrize(
    "mutation",
    ("wrong_owner", "wrong_knowledge_base", "wrong_hash", "not_ready"),
)
def test_s02_v5_source_integrity_anomalies_block_docx(mutation: str) -> None:
    evidence = _snapshot()
    project = evidence[0]
    if mutation == "wrong_owner":
        project["authorization_subject"] = OTHER_OWNER
    elif mutation == "wrong_knowledge_base":
        project["knowledge_base_id"] = "unauthorized-kb"
    elif mutation == "wrong_hash":
        project["content_sha256"] = "0" * 64
    else:
        project["source_status"] = "deleted"

    run, executor, artifact_read, other_user_read = _exercise_s02_v5(evidence)

    assert run.status is RunStatus.FAILED
    assert run.delivery_status is None
    assert run.artifacts == []
    assert executor.artifact_writes == {}
    assert artifact_read is None
    assert other_user_read is None
    assert "来源校验失败" in (run.error or "")
