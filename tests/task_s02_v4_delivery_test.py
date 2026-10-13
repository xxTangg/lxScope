"""S02 v4 delivery modes, server-owned evidence, and artifact gates."""

from __future__ import annotations

import asyncio
import copy
import hashlib
import io
import json
from datetime import datetime, timezone
from typing import Any

from docx import Document
import pytest

from examples.agent_service.task._executor import StepExecutionResult
from examples.agent_service.task._models import (
    AgentStepConfig,
    CreateTaskRequest,
    RunStatus,
    StepType,
    StepVerdict,
    TaskContext,
    TaskInputReply,
    TaskNode,
    TaskRunRequest,
    TaskToolSchema,
    UpdateTaskRequest,
)
from examples.agent_service.task._s02_structured_output import (
    assemble_s02_v4_output,
    assemble_s02_v4_1_output,
)
from examples.agent_service.task._s02_project_planning_evidence import (
    materialize_s02_project_planning_ledger,
)
from examples.agent_service.task._service import TaskService
from examples.agent_service.task import _service as task_service_module
from examples.agent_service.task._store import TaskStore
from examples.agent_service.task._template_compiler import TaskTemplateCompiler
from examples.agent_service.task._template_models import (
    TemplateInstantiateRequest,
)
from examples.agent_service.task._template_registry import TemplateRegistry
from examples.agent_service.task._template_service import TaskTemplateService
from examples.agent_service.task._template_validation import TaskValidationFailure


OWNER = "task-s02-v4-demo-owner"
KB_ID = "kb-s02-v4-demo-fixture"
PROJECT_DOCUMENT_ID = "doc-s02-v4-project-fixture"
TEMPLATE_DOCUMENT_ID = "doc-s02-v4-template-fixture"
PROJECT_QUOTE = "项目名称：[DEMO/SYNTHETIC] 青岚园区策划演示项目。"
TEMPLATE_LIMITATION_QUOTE = (
    "本策划模板为 DEMO/SYNTHETIC 演示材料，不是正式制度或监管要求。"
)
RULE_TEXTS = (
    "明确项目名称、编号和委托对象。",
    "说明咨询范围和不包含事项。",
    "说明能源数据的时间范围和计量单位。",
    "说明表计边界、缺失数据和校准需求。",
    "说明设备台账字段及其来源。",
    "区分已收集资料、用户陈述和待取得资料。",
    "说明访谈对象类别和记录方式，不虚构访谈结果。",
    "说明现场核查步骤和安全前提。",
    "说明分析假设、公式和限制。",
    "区分节能目标建议值与实际节能量。",
    "列出建议措施及其所需验证条件。",
    "对预算数字说明来源和审批状态。",
    "对工期说明关键依赖和不确定性。",
    "列出风险、责任方和需补充证据。",
    "列出交付物、复核人和草稿状态。",
)


def _evidence(
    *,
    evidence_type: str,
    source_text: str,
    source: str,
    document_id: str,
    authorization_subject: str = OWNER,
) -> dict[str, Any]:
    return {
        "knowledge_base_id": KB_ID,
        "knowledge_base_name": "DEMO/SYNTHETIC Task S02 v4 测试知识库",
        "document_id": document_id,
        "source": source,
        "chunk_index": 0,
        "content_sha256": hashlib.sha256(
            source_text.encode("utf-8")
        ).hexdigest(),
        "source_text": source_text,
        "source_text_truncated": False,
        "source_type": "knowledge_base_chunk",
        "source_status": "ready",
        "authorization_subject": authorization_subject,
        "evidence_type": evidence_type,
        "collected_at": datetime.now(timezone.utc).isoformat(),
    }


def _evidence_snapshot() -> list[dict[str, Any]]:
    chapters = (
        "第一章 项目概况",
        "第二章 资料清单与可信度",
        "第三章 能源基线与计量口径",
        "第四章 设备台账现状",
        "第五章 分析方法与工作计划",
        "第六章 方案建议",
        "第七章 预算与进度边界",
        "第八章 风险、待核实事项与交付清单",
    )
    template_text = "\n".join(
        [
            "# 策划报告编制模板",
            "## 八章报告大纲与章节结构",
            *chapters,
            "## B-01 至 B-15 编写要求",
            *(
                f"B-{index:02d} {description}"
                for index, description in enumerate(RULE_TEXTS, 1)
            ),
            TEMPLATE_LIMITATION_QUOTE,
        ]
    )
    project_text = "\n".join(
        [
            "# DEMO/SYNTHETIC 项目委托资料",
            "资料类型：项目委托资料；仅用于 Task 自动化演示。",
            PROJECT_QUOTE,
        ]
    )
    return [
        _evidence(
            evidence_type="project_material",
            source_text=project_text,
            source="DEMO/SYNTHETIC 项目委托资料.md",
            document_id=PROJECT_DOCUMENT_ID,
        ),
        _evidence(
            evidence_type="planning_template",
            source_text=template_text,
            source="DEMO/SYNTHETIC 策划报告编制模板.md",
            document_id=TEMPLATE_DOCUMENT_ID,
        ),
    ]


def _business_payload(*, missing_project_quote: bool = False) -> dict[str, Any]:
    chapters = [{"number": number, "items": []} for number in range(1, 9)]
    chapters[0]["items"].append(
        {
            "kind": "fact",
            "text": PROJECT_QUOTE.removesuffix("。"),
            "source_candidates": (
                ["不存在于授权来源中的伪造摘录"]
                if missing_project_quote
                else [PROJECT_QUOTE]
            ),
        }
    )
    for number in range(1, 9):
        chapters[number - 1]["items"].append(
            {
                "kind": "analysis",
                "text": (
                    f"分析说明：B-{number:02d} 是模板写作编号，"
                    "此处不构成独立规则声明。"
                ),
                "source_candidates": [],
            }
        )
    chapters[5]["items"].append(
        {
            "kind": "recommendation",
            "text": "建议按 B-06 的写作要求区分资料类型，并对缺项保留说明。",
            "source_candidates": [],
        }
    )
    chapters[7]["items"].extend(
        [
            {
                "kind": "template_limitation",
                "text": "策划模板是演示材料，不代表正式制度。",
                "source_candidates": [TEMPLATE_LIMITATION_QUOTE],
            },
            {
                "kind": "unverified",
                "text": "B-06 相关的部分原始项目资料待补充。",
                "source_candidates": [],
                "pending_status": "pending_supplement",
                "reason": "当前授权资料没有该部分原始项目材料。",
                "expected_evidence_types": ["project_material"],
            },
        ]
    )
    return {
        "title": "DEMO/SYNTHETIC 青岚园区项目策划报告",
        "chapters": chapters,
    }


def _assemble(
    *,
    mode: str = "strict_verified",
    payload: dict[str, Any] | None = None,
    evidence: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    template = TemplateRegistry().get_template_version(
        "s02_project_planning", "4.0.0"
    )
    assert template is not None
    contract = template.steps[-1].output_contract
    assert contract is not None and contract.json_schema is not None
    return assemble_s02_v4_output(
        json.dumps(payload or _business_payload(), ensure_ascii=False),
        contract.json_schema,
        _evidence_snapshot() if evidence is None else evidence,
        authorization_subject=OWNER,
        delivery_mode=mode,
    )


def _business_payload_v41(
    *,
    missing_project_quote: bool = False,
    unrelated_project_quote: bool = False,
    composite_project_claim: bool = False,
) -> dict[str, Any]:
    payload = _business_payload(missing_project_quote=missing_project_quote)
    for chapter in payload["chapters"]:
        for item in chapter["items"]:
            if item["kind"] == "unverified":
                continue
            quotes = item["source_candidates"]
            if unrelated_project_quote and item["kind"] == "fact":
                item["text"] = "预算上限 CNY 600,000.00。"
                item["source_candidates"] = [
                    {
                        "quote": "资料类型：项目委托资料；仅用于 Task 自动化演示。",
                        "supports": item["text"],
                    }
                ]
            else:
                item["source_candidates"] = [
                    {"quote": quote, "supports": item["text"]}
                    for quote in quotes
                ]
    if composite_project_claim:
        first = payload["chapters"][0]["items"][0]
        first["text"] = "项目编号为 DEMO-S02-042，建筑面积为 42,000 m²。"
        first["source_candidates"] = [
            {
                "quote": "项目编号：DEMO-S02-042",
                "supports": "项目编号为 DEMO-S02-042",
            },
            {
                "quote": "建筑面积：42,000 m²",
                "supports": "建筑面积为 42,000 m²",
            },
        ]
    return payload


def _assemble_v41(
    *,
    mode: str = "strict_verified",
    payload: dict[str, Any] | None = None,
    evidence: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    template = TemplateRegistry().get_template_version(
        "s02_project_planning", "4.1.0"
    )
    assert template is not None
    contract = template.steps[-1].output_contract
    assert contract is not None and contract.json_schema is not None
    return assemble_s02_v4_1_output(
        json.dumps(
            payload or _business_payload_v41(),
            ensure_ascii=False,
        ),
        contract.json_schema,
        _evidence_snapshot() if evidence is None else evidence,
        authorization_subject=OWNER,
        delivery_mode=mode,
    )
def test_s02_v4_and_v41_historical_templates_keep_versioned_contracts() -> None:
    registry = TemplateRegistry()
    current = registry.get_current_template("s02_project_planning")
    historical_v4 = registry.get_template_version(
        "s02_project_planning", "4.0.0"
    )
    historical_v4_1 = registry.get_template_version(
        "s02_project_planning", "4.1.0"
    )
    assert current is not None and current.template_version == "5.0.0"
    assert historical_v4 is not None and historical_v4_1 is not None
    assert len(current.steps) == 5
    assert all(step.output_contract is None for step in current.steps)
    assert historical_v4.steps[-1].output_contract is not None
    assert historical_v4.steps[-1].output_contract.contract_id == "task.s02.planning.report.v4"
    assert historical_v4_1.steps[-1].output_contract is not None
    assert historical_v4_1.steps[-1].output_contract.contract_id == "task.s02.planning.report.v4.1"
    v4_schema = historical_v4.steps[-1].output_contract.json_schema
    v41_schema = historical_v4_1.steps[-1].output_contract.json_schema
    assert v4_schema is not None and v41_schema is not None
    v4_item = v4_schema["properties"]["chapters"]["items"]["properties"]["items"]["items"]
    v41_item = v41_schema["properties"]["chapters"]["items"]["properties"]["items"]["items"]
    assert v4_item["oneOf"][0]["properties"]["source_candidates"]["items"]["type"] == "string"
    v41_candidate = v41_item["properties"]["source_candidates"]["items"]
    assert v41_candidate["required"] == ["quote", "supports"]
    assert set(v41_candidate["properties"]) == {"quote", "supports"}
    assert "supports" in historical_v4_1.steps[-1].config.prompt

    for version in ("1.0.0", "2.0.0", "3.0.0", "3.1.0", "3.2.0"):
        historical = registry.get_template_version(
            "s02_project_planning", version
        )
        assert historical is not None
        assert historical.template_version == version
    assert current.template_version != historical_v4.template_version
    assert current.template_version != registry.get_template_version(
        "s02_project_planning", "3.2.0"
    ).template_version


def test_s02_v4_strict_assembles_b_reference_text_without_rule_objects() -> None:
    result = _assemble()

    assert result["passed"] is True
    assert result["delivery_report"]["mode"] == "strict_verified"
    assert result["delivery_report"]["chapter_count"] == 8
    assert "B-01 是模板写作编号" in result["document"]
    assert len(result["delivery_report"]["template_rule_requirements"]) == 15
    assert all(
        item["source_available"]
        for item in result["delivery_report"]["template_rule_requirements"]
    )
    assert all(
        claim["kind"] != "rule"
        for claim in json.loads(
            result["document"].split("```task-claim-evidence\n", 1)[1]
            .split("\n```", 1)[0]
        )["claims"]
    )


def _evidence_with_composite_project_facts() -> list[dict[str, Any]]:
    evidence = _evidence_snapshot()
    project = next(
        item for item in evidence if item["evidence_type"] == "project_material"
    )
    project["source_text"] += (
        "\n项目编号：DEMO-S02-042\n建筑面积：42,000 m²"
    )
    project["content_sha256"] = hashlib.sha256(
        project["source_text"].encode("utf-8")
    ).hexdigest()
    return evidence


def test_s02_v41_splits_composite_claim_by_exact_supported_substatements() -> None:
    evidence = _evidence_with_composite_project_facts()
    result = _assemble_v41(
        payload=_business_payload_v41(composite_project_claim=True),
        evidence=evidence,
    )

    assert result["passed"] is True
    assert result["adapter_verification_scope"] == (
        "schema_business_shape_and_source_resolution"
    )
    assert result["delivery_report"]["evidence_ledger_passed"] is False
    protocol = json.loads(
        result["document"].split("```task-claim-evidence\n", 1)[1]
        .split("\n```", 1)[0]
    )
    claims = [
        item
        for item in protocol["claims"]
        if "DEMO-S02-042" in item["text"]
        or "42,000 m²" in item["text"]
    ]
    assert len(claims) == 2
    assert len({item["claim_id"] for item in claims}) == 2
    assert all(item["occurrences"] for item in claims)
    assert all(item["citation_ids"] for item in claims)
    assert len(
        [item for item in protocol["bindings"] if item["claim_id"] in {c["claim_id"] for c in claims}]
    ) == 2
    ledger = materialize_s02_project_planning_ledger(
        result["document"],
        evidence,
        authorization_subject=OWNER,
        template_version="4.1.0",
    )
    assert ledger["passed"] is True


def test_s02_v41_demo_downgrades_unrelated_local_quote_without_forging_ledger() -> None:
    evidence = _evidence_snapshot()
    payload = _business_payload_v41(unrelated_project_quote=True)
    strict = _assemble_v41(payload=payload, evidence=evidence)
    assert strict["passed"] is False
    assert strict["error_code"] == "TASK_S02_V4_CLAIM_QUOTE_ASSOCIATION_FAILED"

    demo = _assemble_v41(
        mode="demo_draft",
        payload=payload,
        evidence=evidence,
    )
    assert demo["passed"] is True
    assert demo["delivery_report"]["status"] == "draft_unverified"
    assert demo["delivery_report"]["evidence_ledger_passed"] is False
    assert any(
        warning["warning_code"] == "quote_unrelated_to_claim"
        and warning["disposition"] == "marked_pending_verification"
        for warning in demo["warnings"]
    )
    protocol = json.loads(
        demo["document"].split("```task-claim-evidence\n", 1)[1]
        .split("\n```", 1)[0]
    )
    downgraded = [
        item
        for item in protocol["claims"]
        if "预算上限 CNY 600,000.00" in item["text"]
    ]
    assert len(downgraded) == 1
    assert downgraded[0]["kind"] == "unverified"
    assert downgraded[0]["citation_ids"] == []
    assert not any(
        binding["claim_id"] == downgraded[0]["claim_id"]
        for binding in protocol["bindings"]
    )
    citation_payload = json.loads(
        demo["document"].split("```task-citations\n", 1)[1]
        .split("\n```", 1)[0]
    )
    assert all(
        "资料类型：项目委托资料；仅用于 Task 自动化演示。" != item["quote"]
        for item in citation_payload["citations"]
    )
    ledger = materialize_s02_project_planning_ledger(
        demo["document"],
        evidence,
        authorization_subject=OWNER,
        template_version="4.1.0",
    )
    assert ledger["passed"] is True


@pytest.mark.parametrize("failure", ["subject", "hash"])
def test_s02_v41_demo_never_downgrades_source_identity_or_hash_failures(
    failure: str,
) -> None:
    evidence = _evidence_snapshot()
    project = next(
        item for item in evidence if item["evidence_type"] == "project_material"
    )
    if failure == "subject":
        project["authorization_subject"] = "another-user"
    else:
        project["content_sha256"] = "0" * 64

    result = _assemble_v41(
        mode="demo_draft",
        payload=_business_payload_v41(),
        evidence=evidence,
    )
    assert result["passed"] is False
    assert result["error_code"] in {
        "TASK_S02_V3_SOURCE_PERMISSION_MISMATCH",
        "TASK_S02_V3_SOURCE_HASH_MISMATCH",
    }


def _reported_run_diagnostic_shape() -> tuple[dict[str, Any], list[dict[str, Any]]]:
    evidence = _evidence_snapshot()
    project = next(
        item for item in evidence if item["evidence_type"] == "project_material"
    )
    unrelated_quotes = [
        f"设备标签记录型号为 QX-{label}。"
        for label in "甲乙丙丁戊己庚"
    ]
    related_quotes = [
        f"授权项目资料载明第{label}类背景信息已经核对。"
        for label in "甲乙丙丁戊己庚辛壬"
    ]
    project["source_text"] += "\n" + "\n".join(
        [*unrelated_quotes, *related_quotes]
    )
    project["content_sha256"] = hashlib.sha256(
        project["source_text"].encode("utf-8")
    ).hexdigest()

    items: list[dict[str, Any]] = []
    for index, quote in enumerate(unrelated_quotes):
        items.append(
            {
                "kind": "fact",
                "text": f"项目预算金额为 {90000 + index} 元并已正式批准。",
                "source_candidates": [quote],
            }
        )
    valid_labels = "甲乙丙丁戊己庚辛壬癸子"
    for index in range(11):
        quote = related_quotes[index % len(related_quotes)]
        items.append(
            {
                "kind": "fact",
                "text": f"{quote}；{valid_labels[index]}事项另有背景说明。",
                "source_candidates": [quote],
            }
        )
    for label in "一二三四五六七":
        items.append(
            {
                "kind": "analysis",
                "text": f"分析意见{label}：工作顺序需要结合资料完整程度调整。",
                "source_candidates": [],
            }
        )
    for index in range(29):
        items.append(
            {
                "kind": "unverified",
                "text": f"待核实事项 DEMO-{index + 1:02d} 的独立资料仍未取得。",
                "source_candidates": [],
                "pending_status": "pending_verification",
                "reason": "当前授权项目资料未包含该事项的核实材料。",
                "expected_evidence_types": ["project_material"],
            }
        )
    assert len(items) == 54
    chapters = []
    cursor = 0
    for number in range(1, 9):
        count = 7 if number <= 6 else 6
        chapters.append({"number": number, "items": items[cursor : cursor + count]})
        cursor += count
    assert cursor == len(items)
    return {"title": "DEMO/SYNTHETIC 诊断形状回归", "chapters": chapters}, evidence


def test_s02_v41_redacted_diagnostic_shape_sanitizes_reported_fourteen_ledger_errors() -> None:
    # The actual Run payload is not in the workspace. This fixture reproduces
    # the supplied counts and diagnostic codes, not the unavailable raw text.
    legacy_payload, evidence = _reported_run_diagnostic_shape()
    legacy = _assemble(
        mode="demo_draft",
        payload=legacy_payload,
        evidence=evidence,
    )
    assert legacy["passed"] is True
    assert (legacy["claim_count"], legacy["occurrence_count"]) == (54, 54)
    assert legacy["binding_count"] == 18
    assert legacy["unverified_count"] == 29
    assert legacy["citation_count"] == 16

    original_ledger = materialize_s02_project_planning_ledger(
        legacy["document"],
        evidence,
        authorization_subject=OWNER,
        template_version="4.0.0",
    )
    codes = [item["error_code"] for item in original_ledger["all_diagnostics"]]
    assert codes.count("quote_unrelated_to_claim") == 7
    assert codes.count("claim_binding_missing") == 7
    assert original_ledger["passed"] is False

    v41_payload = copy.deepcopy(legacy_payload)
    for chapter in v41_payload["chapters"]:
        for item in chapter["items"]:
            if item["kind"] != "unverified":
                item["source_candidates"] = [
                    {"quote": quote, "supports": item["text"]}
                    for quote in item["source_candidates"]
                ]
    sanitized = _assemble_v41(
        mode="demo_draft",
        payload=v41_payload,
        evidence=evidence,
    )
    assert sanitized["passed"] is True
    assert sanitized["raw_model_output"] == json.dumps(
        v41_payload, ensure_ascii=False
    )
    assert len(sanitized["warnings"]) >= 7
    assert len(sanitized["business_result"]["chapters"]) == 8
    sanitized_ledger = materialize_s02_project_planning_ledger(
        sanitized["document"],
        evidence,
        authorization_subject=OWNER,
        template_version="4.1.0",
    )
    assert sanitized_ledger["passed"] is True
    assert sanitized["binding_count"] == 11
    assert sanitized["unverified_count"] == 36


def test_s02_v4_strict_fails_and_demo_marks_unmatched_fact_pending() -> None:
    payload = _business_payload(missing_project_quote=True)
    strict = _assemble(payload=payload)
    demo = _assemble(payload=payload, mode="demo_draft")

    assert strict["passed"] is False
    assert strict["error_code"] == "TASK_S02_V3_SOURCE_QUOTE_NOT_FOUND"
    assert demo["passed"] is True
    assert demo["delivery_report"]["status"] == "draft_unverified"
    assert demo["unverified_count"] == 2
    assert demo["warnings"]
    assert demo["delivery_report"]["human_review_required"] is True
    assert "【待核实】" in demo["document"]
    citations = json.loads(
        demo["document"].split("```task-citations\n", 1)[1]
        .split("\n```", 1)[0]
    )["citations"]
    assert all(item["quote"] != "不存在于授权来源中的伪造摘录" for item in citations)
    normalized = demo["business_result"]
    assert any(
        item["kind"] == "unverified"
        and item["text"].startswith("项目名称")
        for chapter in normalized["chapters"]
        for item in chapter["items"]
    )


def test_s02_v4_demo_never_downgrades_identity_permission_or_hash_failures() -> None:
    unauthorized = _evidence_snapshot()
    unauthorized[0]["authorization_subject"] = "another-user"
    permission_result = _assemble(
        mode="demo_draft", evidence=unauthorized
    )
    assert permission_result["passed"] is False
    assert permission_result["error_code"] == "TASK_S02_V3_SOURCE_PERMISSION_MISMATCH"

    tampered = _evidence_snapshot()
    tampered[0]["source_text"] += "被篡改"
    hash_result = _assemble(mode="demo_draft", evidence=tampered)
    assert hash_result["passed"] is False
    assert hash_result["error_code"] == "TASK_S02_V3_SOURCE_HASH_MISMATCH"


class _NoPlanner:
    async def plan(self, **_: Any) -> Any:
        raise AssertionError("Template instantiation must not invoke Planner.")


class _AuthorizedDemoGateway:
    """Synthetic evidence fixture; it does not represent live RAG success."""

    def __init__(self) -> None:
        snapshot = _evidence_snapshot()
        self.project = [item for item in snapshot if item["evidence_type"] == "project_material"]
        self.template = [item for item in snapshot if item["evidence_type"] == "planning_template"]

    async def retrieve_planning_template_evidence(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        *,
        top_k: int,
    ) -> dict[str, Any]:
        assert user_id == OWNER and knowledge_base_ids == [KB_ID]
        assert top_k == 12
        return {
            "status": "ready",
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "knowledge_bases": [{"id": KB_ID, "name": "DEMO/SYNTHETIC S02 KB"}],
            "evidence": self.template,
            "template_matches": [{"source": self.template[0]["source"]}],
            "queries": ["DEMO/SYNTHETIC planning template"],
            "citation_protocol": "DEMO/SYNTHETIC literal source quote protocol",
        }

    async def retrieve_project_material_evidence(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        *,
        project_name: str,
        source_material_ids: list[str],
        top_k: int,
    ) -> dict[str, Any]:
        assert user_id == OWNER and knowledge_base_ids == [KB_ID]
        assert project_name and source_material_ids == [PROJECT_DOCUMENT_ID]
        assert top_k == 8
        return {
            "status": "ready",
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "knowledge_bases": [{"id": KB_ID, "name": "DEMO/SYNTHETIC S02 KB"}],
            "evidence": self.project,
            "queries": ["DEMO/SYNTHETIC project materials"],
        }

    async def inspect_task_scope(
        self,
        user_id: str,
        knowledge_base_ids: list[str],
        source_material_ids: list[str],
    ) -> dict[str, Any]:
        assert user_id == OWNER and knowledge_base_ids == [KB_ID]
        assert source_material_ids == [PROJECT_DOCUMENT_ID]
        return {
            "knowledge_bases": [{"id": KB_ID, "name": "DEMO/SYNTHETIC S02 KB"}],
        }

    @staticmethod
    def format_task_evidence(result: dict[str, Any]) -> str:
        return json.dumps(result, ensure_ascii=False, default=str)


class _S02V4Executor:
    def __init__(
        self,
        *,
        missing_project_quote: bool,
        fail_artifact_write: bool = False,
        template_version: str = "4.0.0",
        payload: dict[str, Any] | None = None,
    ) -> None:
        if payload is not None:
            business_payload = payload
        elif template_version == "4.1.0":
            business_payload = _business_payload_v41(
                missing_project_quote=missing_project_quote
            )
        else:
            business_payload = _business_payload(
                missing_project_quote=missing_project_quote
            )
        self.raw_output = json.dumps(business_payload, ensure_ascii=False)
        self.artifact_writes: dict[str, bytes] = {}
        self.fail_artifact_write = fail_artifact_write

    async def list_tools(self, _context: TaskContext) -> list[TaskToolSchema]:
        return []

    async def execute_node(
        self,
        *,
        node: Any,
        previous_output: str,
        context: Any,
    ) -> StepExecutionResult:
        del previous_output, context
        if node.name == "整理策划输出":
            return StepExecutionResult(
                text=self.raw_output,
                result={
                    "status": "completed",
                    "output": self.raw_output,
                    "reason": "",
                    "questions": [],
                },
            )
        return StepExecutionResult(text="DEMO/SYNTHETIC 已处理上游内容。")

    async def verify_node(self, **_: Any) -> StepVerdict:
        return StepVerdict(passed=True)

    async def preflight_artifact_writer(
        self,
        *,
        context: Any,
    ) -> tuple[bool, str]:
        assert context.user_id == OWNER
        return True, "DEMO/SYNTHETIC Workspace fixture available"

    async def write_artifact(
        self,
        *,
        path: str,
        data: bytes,
        context: Any,
    ) -> int:
        assert context.user_id == OWNER
        if self.fail_artifact_write:
            raise OSError("DEMO/SYNTHETIC Workspace write failure")
        self.artifact_writes[path] = data
        return len(data)

    async def read_artifact(self, *, path: str, context: Any) -> bytes:
        assert context.user_id == OWNER
        return self.artifact_writes[path]


async def _execute_v4(
    mode: str,
    *,
    missing_project_quote: bool,
    fail_artifact_write: bool = False,
    template_version: str = "4.0.0",
    payload: dict[str, Any] | None = None,
) -> tuple[Any, _S02V4Executor]:
    executor = _S02V4Executor(
        missing_project_quote=missing_project_quote,
        fail_artifact_write=fail_artifact_write,
        template_version=template_version,
        payload=payload,
    )
    service = TaskService(
        TaskStore(),
        agentscope_executor=executor,  # type: ignore[arg-type]
        preview_executor=executor,  # type: ignore[arg-type]
        planner=_NoPlanner(),  # type: ignore[arg-type]
        artifact_writer=executor,  # type: ignore[arg-type]
    )
    gateway = _AuthorizedDemoGateway()
    try:
        task = await TaskTemplateService(
            TemplateRegistry(), service, compiler=TaskTemplateCompiler()
        ).instantiate(
            OWNER,
            "s02_project_planning",
            TemplateInstantiateRequest(
                template_version=template_version,
                inputs={
                    "project_name": "DEMO/SYNTHETIC 青岚园区策划演示项目",
                    "project_brief": "仅用于验证 S02 v4 输出与产物门控。",
                    "source_material_ids": [PROJECT_DOCUMENT_ID],
                },
                source_context=TaskContext(
                    agent_id="demo-s02-v4-agent",
                    session_id="demo-s02-v4-session",
                ),
            ),
        )
        task = await service.update_task(
            OWNER,
            task.id,
            UpdateTaskRequest(knowledge_base_ids=[KB_ID]),
        )
        run = await service.create_run(
            OWNER,
            task.id,
            TaskRunRequest(delivery_mode=mode),
            gateway,  # type: ignore[arg-type]
        )
        assert run is not None
        for _ in range(10):
            active = service._active_runs.get(run.id)
            if active is not None:
                await active
            run = await service.get_run(OWNER, run.id)
            assert run is not None
            if run.status in {RunStatus.SUCCEEDED, RunStatus.FAILED}:
                return run, executor
            assert run.status is RunStatus.AWAITING_INPUT
            assert run.pending_input is not None
            answers: dict[str, str] = {}
            for question in run.pending_input.questions:
                if question.key == "project_name":
                    answers[question.key] = "DEMO/SYNTHETIC 青岚园区策划演示项目"
                elif question.key == "project_brief":
                    answers[question.key] = "仅用于验证 S02 v4 输出与产物门控。"
                elif question.key == "source_material_ids":
                    answers[question.key] = json.dumps([PROJECT_DOCUMENT_ID])
                elif question.key == "planning_feedback":
                    answers[question.key] = "无补充"
                elif question.required:
                    answers[question.key] = "DEMO/SYNTHETIC 测试输入"
            await service.answer_run(
                OWNER,
                run.id,
                TaskInputReply(
                    request_id=run.pending_input.id,
                    answers=answers,
                ),
                gateway,  # type: ignore[arg-type]
            )
        raise AssertionError("S02 v4 test Run did not reach a terminal state.")
    finally:
        await service.shutdown()


def test_s02_v4_taskservice_sop_demo_draft_writes_marked_docx_after_ledger() -> None:
    run, executor = asyncio.run(
        _execute_v4("demo_draft", missing_project_quote=True)
    )

    assert run.status is RunStatus.SUCCEEDED
    assert run.source_template_version == "4.0.0"
    assert run.delivery_mode == "demo_draft"
    assert run.delivery_status == "draft_unverified"
    assert run.verification_report["evidence_ledger_passed"] is True
    assert run.verification_report["chapter_count"] == 8
    assert run.verification_report["source_ledger"]
    assert len(run.artifacts) == 1
    artifact = run.artifacts[0]
    assert artifact.run_id == run.id
    assert artifact.delivery_status == "draft_unverified"
    assert len(executor.artifact_writes) == 1
    docx_bytes = next(iter(executor.artifact_writes.values()))
    assert docx_bytes.startswith(b"PK")
    document = Document(io.BytesIO(docx_bytes))
    all_text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert "DEMO/SYNTHETIC · 演示草稿 · 未完成完整证据核验" in all_text
    assert "draft_unverified" in document.core_properties.subject
    final_node = run.node_runs[-1]
    assert final_node.metadata["evidence_ledger"]["passed"] is True
    assert final_node.metadata["s02_v4_delivery"]["passed"] is True
    assert any(
        warning["disposition"] == "marked_pending_verification"
        for warning in final_node.metadata["s02_v4_delivery"]["delivery_report"][
            "warnings"
        ]
    )


def test_s02_v41_taskservice_sop_demo_writes_marked_docx_after_safe_downgrade() -> None:
    run, executor = asyncio.run(
        _execute_v4(
            "demo_draft",
            missing_project_quote=False,
            template_version="4.1.0",
            payload=_business_payload_v41(unrelated_project_quote=True),
        )
    )

    assert run.status is RunStatus.SUCCEEDED
    assert run.source_template_version == "4.1.0"
    assert run.delivery_status == "draft_unverified"
    assert len(run.artifacts) == 1
    artifact = run.artifacts[0]
    assert artifact.run_id == run.id
    assert artifact.delivery_status == "draft_unverified"
    assert len(executor.artifact_writes) == 1
    docx_bytes = next(iter(executor.artifact_writes.values()))
    assert docx_bytes.startswith(b"PK")
    docx = Document(io.BytesIO(docx_bytes))
    content = "\n".join(paragraph.text for paragraph in docx.paragraphs)
    assert "DEMO/SYNTHETIC · 演示草稿 · 未完成完整证据核验" in content
    assert docx.paragraphs[0].text == (
        "DEMO/SYNTHETIC · 演示草稿 · 未完成完整证据核验"
    )
    assert any(
        "DEMO/SYNTHETIC · 演示草稿 · 未完成完整证据核验"
        in paragraph.text
        for section in docx.sections
        for paragraph in section.header.paragraphs
    )
    final_node = run.node_runs[-1]
    assert final_node.metadata["s02_v4_delivery"]["passed"] is True
    assert final_node.metadata["evidence_ledger"]["passed"] is True
    assert final_node.metadata["output_contract_validation"]["target"] == (
        "s02_v4_1_raw_agent_json"
    )
    assert any(
        warning["warning_code"] == "quote_unrelated_to_claim"
        for warning in final_node.metadata["s02_v4_delivery"]["delivery_report"][
            "warnings"
        ]
    )
    assert final_node.metadata["evidence_ledger"]["verification_scope"][
        "claim_quote_association"
    ] == "passed"


def test_s02_v41_adapter_pass_does_not_override_ledger_failure_or_write_docx(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_ledger(*_: Any, **__: Any) -> dict[str, Any]:
        return {
            "passed": False,
            "issues": ["DEMO/SYNTHETIC forced independent ledger failure"],
            "diagnostics": {"citation": []},
            "all_diagnostics": [],
            "ledger": [],
            "verification_scope": {
                "authorized_source_identity_and_hash": "passed",
                "claim_quote_association": "failed",
            },
        }

    monkeypatch.setattr(
        task_service_module,
        "materialize_s02_project_planning_ledger",
        fail_ledger,
    )
    run, executor = asyncio.run(
        _execute_v4(
            "demo_draft",
            missing_project_quote=False,
            template_version="4.1.0",
            payload=_business_payload_v41(),
        )
    )

    assert run.status is RunStatus.FAILED
    assert run.artifacts == []
    assert executor.artifact_writes == {}
    final_node = run.node_runs[-1]
    assert final_node.metadata["s02_v4_delivery"]["passed"] is True
    assert final_node.metadata["evidence_ledger"]["passed"] is False
    assert (
        final_node.metadata["evidence_ledger"]["verification_scope"][
            "claim_quote_association"
        ]
        == "failed"
    )


def test_s02_v41_strict_association_failure_blocks_taskrun_and_docx() -> None:
    run, executor = asyncio.run(
        _execute_v4(
            "strict_verified",
            missing_project_quote=False,
            template_version="4.1.0",
            payload=_business_payload_v41(unrelated_project_quote=True),
        )
    )

    assert run.status is RunStatus.FAILED
    assert run.delivery_status is None
    assert run.artifacts == []
    assert executor.artifact_writes == {}


def test_s02_v4_strict_missing_source_blocks_success_status_and_docx() -> None:
    run, executor = asyncio.run(
        _execute_v4("strict_verified", missing_project_quote=True)
    )

    assert run.status is RunStatus.FAILED
    assert run.delivery_status is None
    assert run.verification_report == {}
    assert run.artifacts == []
    assert executor.artifact_writes == {}


def test_s02_v4_workspace_write_failure_does_not_mark_delivery_or_artifact() -> None:
    run, executor = asyncio.run(
        _execute_v4(
            "demo_draft",
            missing_project_quote=False,
            fail_artifact_write=True,
        )
    )

    assert run.status is RunStatus.FAILED
    assert run.delivery_status is None
    assert run.verification_report == {}
    assert run.artifacts == []
    assert executor.artifact_writes == {}


def test_demo_mode_is_rejected_for_free_planner_tasks() -> None:
    async def exercise() -> None:
        service = TaskService(TaskStore(), planner=_NoPlanner())  # type: ignore[arg-type]
        try:
            task = await service.create_task(
                OWNER,
                CreateTaskRequest(
                    goal="DEMO/SYNTHETIC mode restriction check",
                    nodes=[
                        TaskNode(
                            name="普通自由任务步骤",
                            type=StepType.AGENT,
                            config=AgentStepConfig(prompt="处理任务"),
                        )
                    ],
                ),
            )
            with pytest.raises(TaskValidationFailure):
                await service.create_run(
                    OWNER,
                    task.id,
                    TaskRunRequest(delivery_mode="demo_draft"),
                )
        finally:
            await service.shutdown()

    asyncio.run(exercise())
