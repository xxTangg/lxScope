# -*- coding: utf-8 -*-
"""Format-aware Office renderers owned by the Task module."""

from __future__ import annotations

import datetime as dt
import io
import re
from typing import Any

from ._artifact_spec import ArtifactSpec, ArtifactTable, spec_from_markdown


def _coerce_spec(value: ArtifactSpec | str) -> ArtifactSpec:
    return spec_from_markdown(value) if isinstance(value, str) else value


def _is_weather_table(table: ArtifactTable) -> bool:
    columns = " ".join(table.columns)
    return "日期" in columns and ("最高" in columns or "最低" in columns or "气温" in columns)


def _weather_tables(table: ArtifactTable) -> tuple[ArtifactTable, ...]:
    """Split wide weather records from long recommendations."""

    recommendation_indexes = tuple(
        index
        for index, column in enumerate(table.columns)
        if any(word in column for word in ("穿衣", "出行", "建议", "提示"))
    )
    date_indexes = tuple(index for index, column in enumerate(table.columns) if "日期" in column)
    if not recommendation_indexes or not date_indexes:
        return (table,)
    core_indexes = tuple(index for index in range(len(table.columns)) if index not in recommendation_indexes)
    advice_indexes = date_indexes + recommendation_indexes
    core = ArtifactTable(
        title="天气明细",
        columns=tuple(table.columns[index] for index in core_indexes),
        rows=tuple(tuple(row[index] if index < len(row) else "" for index in core_indexes) for row in table.rows),
    )
    advice = ArtifactTable(
        title="出行建议",
        columns=tuple(table.columns[index] for index in advice_indexes),
        rows=tuple(tuple(row[index] if index < len(row) else "" for index in advice_indexes) for row in table.rows),
    )
    return core, advice


def _expanded_tables(spec: ArtifactSpec) -> list[ArtifactTable]:
    result: list[ArtifactTable] = []
    for table in spec.tables:
        result.extend(_weather_tables(table) if _is_weather_table(table) else (table,))
    return result


def _set_cell_margins(cell: Any, qn: Any, OxmlElement: Any) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side in ("top", "start", "bottom", "end"):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), "90")
        element.set(qn("w:type"), "dxa")


def _word_table(document: Any, table_spec: ArtifactTable) -> Any:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    columns = max(1, len(table_spec.columns))
    table = document.add_table(rows=len(table_spec.rows) + 1, cols=columns)
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    header_properties = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    values = [table_spec.columns] + [row for row in table_spec.rows]
    long_column = lambda value: any(word in value for word in ("建议", "提示", "说明", "描述"))
    if _is_weather_table(table_spec):
        widths = [2.2 if "日期" in column else 1.8 if any(word in column for word in ("温度", "天气", "风力")) else 2.2 for column in table_spec.columns]
    else:
        base = 16.0 / columns
        widths = [max(2.0, min(6.0, 5.2 if long_column(column) else base)) for column in table_spec.columns]
        total = sum(widths)
        widths = [width * 16.0 / total for width in widths]
    for row_index, row_values in enumerate(values):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.text = str(row_values[column_index]) if column_index < len(row_values) else ""
            cell.width = Cm(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, qn, OxmlElement)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "4472C4")
                cell._tc.get_or_add_tcPr().append(shading)
            elif row_index % 2 == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EAF2F8")
                cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(9 if columns >= 7 else 9.5)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    return table


def render_docx(value: ArtifactSpec | str) -> bytes:
    """Render a semantic ArtifactSpec as a readable Word report."""

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    spec = _coerce_spec(value)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    for name, size in (("Title", 20), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    title = document.add_paragraph(spec.title, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if spec.subtitle:
        subtitle = document.add_paragraph(spec.subtitle)
        subtitle.paragraph_format.space_after = Pt(10)
        for run in subtitle.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(89, 89, 89)

    if spec.summary:
        document.add_heading("摘要", level=1)
        for label, value_text in spec.summary:
            paragraph = document.add_paragraph()
            lead = paragraph.add_run(f"{label}：")
            lead.bold = True
            paragraph.add_run(value_text)

    rendered_tables: set[int] = set()
    for section_spec in spec.sections:
        document.add_heading(section_spec.title, level=1)
        for paragraph_text in section_spec.paragraphs:
            document.add_paragraph(paragraph_text)
        for bullet in section_spec.bullets:
            document.add_paragraph(bullet, style="List Bullet")
        for table_index in section_spec.table_indexes:
            if 0 <= table_index < len(spec.tables):
                source = spec.tables[table_index]
                for table in _weather_tables(source) if _is_weather_table(source) else (source,):
                    document.add_heading(table.title, level=2)
                    _word_table(document, table)
                    document.add_paragraph()
                rendered_tables.add(table_index)
    for index, source in enumerate(spec.tables):
        if index in rendered_tables:
            continue
        for table in _weather_tables(source) if _is_weather_table(source) else (source,):
            document.add_heading(table.title, level=1)
            _word_table(document, table)
            document.add_paragraph()
    if spec.notes:
        document.add_heading("说明", level=1)
        for note in spec.notes:
            document.add_paragraph(note, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _safe_sheet_name(value: str, used: set[str], index: int) -> str:
    candidate = re.sub(r"[\\/:?*\[\]]+", " ", value).strip() or f"数据表{index}"
    candidate = candidate[:31]
    suffix = 1
    result = candidate
    while result in used:
        suffix += 1
        marker = f"_{suffix}"
        result = f"{candidate[:31 - len(marker)]}{marker}"
    used.add(result)
    return result


def _typed_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        try:
            return dt.date.fromisoformat(text)
        except ValueError:
            return text
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+\.\d+", text):
        return float(text)
    return value


def _unique_headers(headers: tuple[str, ...]) -> list[str]:
    result: list[str] = []
    used: set[str] = set()
    for index, header in enumerate(headers, start=1):
        base = header.strip() or f"列{index}"
        candidate = base
        suffix = 1
        while candidate in used:
            suffix += 1
            candidate = f"{base}_{suffix}"
        used.add(candidate)
        result.append(candidate)
    return result


def _write_xlsx_table(workbook: Any, table_spec: ArtifactTable, sheet_name: str, table_number: int) -> Any:
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(sheet_name)
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = table_spec.title
    sheet["A1"].font = Font(name="Microsoft YaHei", size=15, bold=True, color="000000")
    sheet["A2"] = "可筛选明细"
    sheet["A2"].font = Font(name="Microsoft YaHei", italic=True, color="666666")
    headers = _unique_headers(table_spec.columns)
    start_row = 4
    for column_index, header in enumerate(headers, start=1):
        sheet.cell(start_row, column_index, header)
    for row_index, row in enumerate(table_spec.rows, start=start_row + 1):
        for column_index, value in enumerate(row, start=1):
            sheet.cell(row_index, column_index, _typed_value(value))
    last_column = max(1, len(headers))
    last_row = max(start_row, start_row + len(table_spec.rows))
    ref = f"A{start_row}:{get_column_letter(last_column)}{last_row}"
    table = Table(displayName=f"TaskDataTable{table_number}", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = f"A{start_row + 1}"
    for row in sheet.iter_rows(min_row=start_row, max_row=last_row, max_col=last_column):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=cell.row == start_row, color="FFFFFF" if cell.row == start_row else "000000")
            if cell.row == start_row:
                cell.fill = PatternFill("solid", fgColor="4472C4")
            if isinstance(cell.value, dt.date):
                cell.number_format = "yyyy-mm-dd"
    for column_index, header in enumerate(headers, start=1):
        values = [str(sheet.cell(row, column_index).value or "") for row in range(start_row, last_row + 1)]
        if any(word in header for word in ("建议", "提示", "说明", "描述")):
            width = 34
        elif "日期" in header:
            width = 13
        else:
            width = min(max(max((len(value) for value in values), default=10) + 2, 10), 20)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    for row_index in range(start_row + 1, last_row + 1):
        if any(any(word in header for word in ("建议", "提示", "说明", "描述")) for header in headers):
            sheet.row_dimensions[row_index].height = 42
    if _is_weather_table(table_spec):
        high = next((index + 1 for index, header in enumerate(headers) if "最高" in header), None)
        low = next((index + 1 for index, header in enumerate(headers) if "最低" in header), None)
        date_column = next((index + 1 for index, header in enumerate(headers) if "日期" in header), None)
        if high and low and date_column and len(table_spec.rows) >= 2:
            chart = LineChart()
            chart.title = "气温变化"
            chart.y_axis.title = "温度 ℃"
            chart.x_axis.title = "日期"
            chart.height = 7
            chart.width = 13
            data = Reference(sheet, min_col=high, max_col=low, min_row=start_row, max_row=last_row)
            categories = Reference(sheet, min_col=date_column, min_row=start_row + 1, max_row=last_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            sheet.add_chart(chart, f"{get_column_letter(last_column + 2)}{start_row}")
    return sheet


def render_xlsx(value: ArtifactSpec | str) -> bytes:
    """Render a semantic ArtifactSpec as a reader-oriented workbook."""

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    spec = _coerce_spec(value)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "概览"
    overview.sheet_view.showGridLines = False
    overview["A1"] = spec.title
    overview["A1"].font = Font(name="Microsoft YaHei", size=17, bold=True, color="000000")
    if spec.subtitle:
        overview["A2"] = spec.subtitle
        overview["A2"].font = Font(name="Microsoft YaHei", italic=True, color="666666")
    overview["A4"] = "关键摘要"
    overview["A4"].font = Font(name="Microsoft YaHei", bold=True, color="1F4E78")
    overview["A4"].fill = PatternFill("solid", fgColor="D9EAF7")
    row = 5
    for label, summary in spec.summary:
        overview.cell(row, 1, label)
        overview.cell(row, 2, summary)
        row += 1
    if spec.summary:
        row += 1
    overview.cell(row, 1, "内容导航")
    overview.cell(row, 1).font = Font(name="Microsoft YaHei", bold=True, color="1F4E78")
    overview.cell(row, 1).fill = PatternFill("solid", fgColor="D9EAF7")
    row += 1
    expanded = _expanded_tables(spec)
    used_names = {"概览"}
    table_names: list[str] = []
    table_number = 1
    for table in expanded:
        name = _safe_sheet_name(table.title, used_names, table_number)
        table_names.append(name)
        table_number += 1
    for section in spec.sections:
        overview.cell(row, 1, section.title)
        overview.cell(row, 1).font = Font(name="Microsoft YaHei", bold=True)
        for paragraph in section.paragraphs:
            row += 1
            overview.cell(row, 2, paragraph)
        for bullet in section.bullets:
            row += 1
            overview.cell(row, 2, f"• {bullet}")
        row += 1
    for name in table_names:
        overview.cell(row, 1, "数据表")
        overview.cell(row, 2, name)
        row += 1
    overview.column_dimensions["A"].width = 18
    overview.column_dimensions["B"].width = 72
    for current_row in overview.iter_rows(min_row=1, max_row=max(1, row - 1), max_col=2):
        for cell in current_row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    overview.freeze_panes = "A5" if spec.summary else None

    for index, (table, sheet_name) in enumerate(zip(expanded, table_names), start=1):
        _write_xlsx_table(workbook, table, sheet_name, index)
    if spec.notes:
        notes = workbook.create_sheet(_safe_sheet_name("说明", used_names, table_number))
        notes.sheet_view.showGridLines = False
        notes["A1"] = "说明"
        notes["A1"].font = Font(name="Microsoft YaHei", size=15, bold=True)
        notes.column_dimensions["A"].width = 100
        for index, note in enumerate(spec.notes, start=3):
            notes.cell(index, 1, note)
            notes.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="top")
            notes.row_dimensions[index].height = 36

    workbook.properties.title = spec.title
    workbook.properties.subject = "AgentScope Task 结构化文件产物"
    workbook.properties.creator = "AgentScope Task"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
